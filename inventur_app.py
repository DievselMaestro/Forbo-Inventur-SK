#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
INVENTORY PROGRAM FOR FORBO - WAREHOUSE MANAGEMENT
==================================================

Desktop application for warehouse inventory with barcode scanner integration.
Warehouse modes: HALB, ZERT, KMAT, WIP, Assembling, Service, Warehouse, Granules.
Assembling, Service and Warehouse behave like HALB (rolls).
Developed for Windows 11, Python 3.11+

Date: October 2026
"""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from datetime import datetime
import os
import sys
import shutil
from pathlib import Path
import json
import logging


APP_VERSION = "3.0"

# HALB-type modes share the roll workflow (QR + Shelf Location + Measured Width).
# Value = config key suffix: arbeitstabelle{suffix}_path / export{suffix}_path.
# Output file: Inventory_<mode>.xlsx
HALB_MODES = {
    "HALB": "",
    "Assembling": "_assembling",
    "Service": "_service",
    "Warehouse": "_warehouse",
}

# Mode id -> (display name, colour). Order = order of buttons and settings tabs.
MODES = {
    "HALB": ("HALB", "#1f4e79"),
    "Zert": ("ZERT", "#375623"),
    "KMAT": ("KMAT", "#6b1f1f"),
    "WIP": ("WIP", "#4a3728"),
    "Assembling": ("ASSEMBLING", "#2f5f8f"),
    "Service": ("SERVICE", "#5b3f7a"),
    "Warehouse": ("WAREHOUSE", "#1f6b6b"),
    "Granules": ("GRANULES", "#8a5a00"),
}

# Mode id -> (master table config key, export folder config key)
MODE_CONFIG_KEYS = {
    **{m: (f"arbeitstabelle{s}_path", f"export{s}_path") for m, s in HALB_MODES.items()},
    "Zert": ("arbeitstabelle_zert_path", "export_zert_path"),
    "KMAT": ("arbeitstabelle_kmat_path", "export_kmat_path"),
    "WIP": ("arbeitstabelle_wip_path", "export_wip_path"),
    "Granules": ("arbeitstabelle_granules_path", "export_granules_path"),
}

# Granules master table: SAP export may come with German or English headers.
# Maps every accepted header to the internal (English) column name.
GRANULES_COLUMN_MAP = {
    "Werk": "Plant",
    "Plant": "Plant",
    "LOrt": "Storage Location",
    "Storage Location": "Storage Location",
    "Materialnummer": "Material Number",
    "Material": "Material Number",
    "Material Number": "Material Number",
    "Materialkurztext": "Material Description",
    "Material Description": "Material Description",
    "Charge": "Batch",
    "Batch": "Batch",
    "Frei verwendbar": "Unrestricted",
    "Unrestricted": "Unrestricted",
    "BME": "Base Unit of Measure",
    "Base Unit of Measure": "Base Unit of Measure",
}


# ---------------------------------------------------------------------------
# Shared scan-resolution helpers (WIP / KMAT)
# ---------------------------------------------------------------------------

def _norm_num(x):
    """Normalize a numeric-ish value to a bare digit string, no leading zeros."""
    if x is None:
        return ""
    s = str(x).strip()
    if s in ("", "nan"):
        return ""
    try:
        return str(int(float(s)))
    except (ValueError, TypeError):
        return ""


def _norm_num_lookup(x):
    """Like _norm_num, but keeps the stripped original on failure instead of
    "" so a malformed query value still fails to match rather than matching
    blank cells."""
    s = str(x).strip() if x is not None else ""
    try:
        return str(int(float(s)))
    except (ValueError, TypeError):
        return s


def _batch_key(x):
    """Granules batch key: stripped, leading zeros ignored (0618604585 == 618604585)."""
    s = str(x).strip() if x is not None else ""
    if s.lower() == "nan":
        return ""
    return s.lstrip("0") or s


def _to_float(x):
    """Parse a quantity (accepts decimal comma). Returns None if not numeric."""
    try:
        return float(str(x).strip().replace(",", "."))
    except (ValueError, TypeError):
        return None


def _fmt_qty(x):
    """Format a numeric value without trailing .0; non-numeric values unchanged."""
    f = _to_float(x)
    if f is None:
        return "" if x is None or str(x).lower() == "nan" else str(x)
    return str(int(f)) if f.is_integer() else f"{f:g}"


# ---------------------------------------------------------------------------
# Warehouse Selection Dialog
# ---------------------------------------------------------------------------

class WarehouseSelectionDialog:
    """Startup dialog: select which warehouse to scan."""

    def __init__(self, parent):
        self.result = None

        self.dialog = tk.Toplevel(parent)
        self.dialog.title(f"Select Warehouse - Inventory Forbo v{APP_VERSION}")
        self.dialog.geometry("900x340")
        self.dialog.resizable(False, False)
        self.dialog.grab_set()
        # Center on screen
        self.dialog.update_idletasks()
        x = (self.dialog.winfo_screenwidth() // 2) - 450
        y = (self.dialog.winfo_screenheight() // 2) - 170
        self.dialog.geometry(f"+{x}+{y}")
        self.dialog.protocol("WM_DELETE_WINDOW", self._cancel)

        self._build_widgets()
        self.dialog.wait_window()

    def _build_widgets(self):
        frame = ttk.Frame(self.dialog, padding="20")
        frame.pack(fill=tk.BOTH, expand=True)

        ttk.Label(
            frame,
            text="Which warehouse to scan?",
            font=("Arial", 13, "bold"),
            justify=tk.CENTER,
        ).pack(pady=(0, 20))

        # 4 buttons per row
        mode_ids = list(MODES)
        for start in range(0, len(mode_ids), 4):
            row_frame = ttk.Frame(frame)
            row_frame.pack(pady=(0, 10))
            for i, mode in enumerate(mode_ids[start:start + 4]):
                label, color = MODES[mode]
                tk.Button(
                    row_frame,
                    text=label,
                    font=("Arial", 14, "bold"),
                    bg=color,
                    fg="white",
                    width=14,
                    height=2,
                    command=lambda m=mode: self._select(m),
                ).pack(side=tk.LEFT, padx=(0 if i == 0 else 20, 0))

        ttk.Label(frame, text=f"Version {APP_VERSION}", font=("Arial", 9),
                  foreground="gray").pack(side=tk.BOTTOM, anchor=tk.E)

    def _select(self, mode):
        self.result = mode
        self.dialog.destroy()

    def _cancel(self):
        self.result = None
        self.dialog.destroy()


# ---------------------------------------------------------------------------
# Settings Dialog
# ---------------------------------------------------------------------------

class SettingsDialog:
    """Dialog for configuring file paths (master table + export folder) per mode.
    One tab per warehouse mode; the tab of the active mode is selected."""

    def __init__(self, parent, current_config, current_mode=None):
        self.result = None
        self.config = current_config.copy()
        self._path_vars = {}  # config key -> StringVar

        self.dialog = tk.Toplevel(parent)
        self.dialog.title("Settings")
        self.dialog.geometry("660x360")
        self.dialog.resizable(True, False)
        self.dialog.transient(parent)
        self.dialog.grab_set()
        self.dialog.geometry(
            "+%d+%d" % (parent.winfo_rootx() + 60, parent.winfo_rooty() + 60)
        )

        self._build_widgets(current_mode)
        self.dialog.wait_window()

    def _build_widgets(self, current_mode):
        frame = ttk.Frame(self.dialog, padding="15")
        frame.pack(fill=tk.BOTH, expand=True)
        frame.columnconfigure(0, weight=1)

        notebook = ttk.Notebook(frame)
        notebook.grid(row=0, column=0, sticky=(tk.W, tk.E))

        for mode, (label, _color) in MODES.items():
            master_key, export_key = MODE_CONFIG_KEYS[mode]
            tab = ttk.Frame(notebook, padding="12")
            tab.columnconfigure(0, weight=1)
            notebook.add(tab, text=label)
            if mode == current_mode:
                notebook.select(tab)

            # --- Master Table path ---
            ttk.Label(tab, text=f"{label} Master Table file (*.xlsx):",
                      font=("Arial", 10, "bold")).grid(row=0, column=0, sticky=tk.W, pady=(0, 4))
            master_var = tk.StringVar(value=self.config.get(master_key, ""))
            ttk.Entry(tab, textvariable=master_var, width=55).grid(
                row=1, column=0, sticky=(tk.W, tk.E), padx=(0, 8))
            ttk.Button(tab, text="Browse...",
                       command=lambda v=master_var, l=label: self._browse_file(v, l)).grid(
                row=1, column=1, sticky=tk.W)
            self._path_vars[master_key] = master_var

            # --- Export folder ---
            ttk.Label(tab, text=f"{label} Export output folder:",
                      font=("Arial", 10, "bold")).grid(row=2, column=0, sticky=tk.W, pady=(16, 4))
            export_var = tk.StringVar(value=self.config.get(export_key, ""))
            ttk.Entry(tab, textvariable=export_var, width=55).grid(
                row=3, column=0, sticky=(tk.W, tk.E), padx=(0, 8))
            ttk.Button(tab, text="Browse...",
                       command=lambda v=export_var, l=label: self._browse_dir(v, l)).grid(
                row=3, column=1, sticky=tk.W)
            self._path_vars[export_key] = export_var

        # --- Switch Warehouse button toggle ---
        self.switch_warehouse_var = tk.BooleanVar(
            value=self.config.get("show_switch_warehouse", False))
        ttk.Checkbutton(
            frame,
            text="Show 'Switch Warehouse' button (for presentations)",
            variable=self.switch_warehouse_var,
        ).grid(row=1, column=0, sticky=tk.W, pady=(16, 0))

        # --- Buttons ---
        btn_frame = ttk.Frame(frame)
        btn_frame.grid(row=2, column=0, pady=(16, 0))

        ttk.Button(btn_frame, text="Save", command=self._save, width=12).pack(
            side=tk.LEFT, padx=(0, 12))
        ttk.Button(btn_frame, text="Cancel", command=self._cancel, width=12).pack(
            side=tk.LEFT)

        ttk.Label(frame, text=f"Version {APP_VERSION}", font=("Arial", 9),
                  foreground="gray").grid(row=3, column=0, sticky=tk.E, pady=(8, 0))

        self.dialog.bind("<Return>", lambda e: self._save())
        self.dialog.bind("<Escape>", lambda e: self._cancel())

    def _browse_file(self, var, label):
        path = filedialog.askopenfilename(
            title=f"Select {label} Master Table file",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
        )
        if path:
            var.set(path)

    def _browse_dir(self, var, label):
        path = filedialog.askdirectory(title=f"Select {label} Export output folder")
        if path:
            var.set(path)

    def _save(self):
        for key, var in self._path_vars.items():
            self.config[key] = var.get().strip()
        self.config["show_switch_warehouse"] = self.switch_warehouse_var.get()
        self.result = self.config
        self.dialog.destroy()

    def _cancel(self):
        self.result = None
        self.dialog.destroy()


# ---------------------------------------------------------------------------
# NotFoundDialogHALB
# ---------------------------------------------------------------------------

class NotFoundDialogHALB:
    """Dialog for rolls not found in the master table.
    Pre-fills Batch No., Location and all QR dimensions; user supplies Material,
    Description, Shelf Location (mandatory) and Remarks (optional).
    Always treated as a roll — no type selector.
    """

    def __init__(self, parent, charge, qr_data):
        self.result = None

        self.dialog = tk.Toplevel(parent)
        self.dialog.title("Roll Not Found - Manual Entry")
        self.dialog.geometry("860x820")
        self.dialog.resizable(True, True)
        self.dialog.transient(parent)
        self.dialog.grab_set()
        self.dialog.geometry(
            "+%d+%d" % (parent.winfo_rootx() + 60, parent.winfo_rooty() + 60)
        )

        self._qr = qr_data  # dict from parse_qr_code
        self._build_widgets(charge)
        self.dialog.wait_window()

    def _build_widgets(self, charge):
        canvas = tk.Canvas(self.dialog, borderwidth=0)
        scrollbar = ttk.Scrollbar(self.dialog, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        inner = ttk.Frame(canvas, padding="15")
        win_id = canvas.create_window((0, 0), window=inner, anchor="nw")

        def _on_frame_configure(event):
            canvas.configure(scrollregion=canvas.bbox("all"))

        def _on_canvas_configure(event):
            canvas.itemconfig(win_id, width=event.width)

        inner.bind("<Configure>", _on_frame_configure)
        canvas.bind("<Configure>", _on_canvas_configure)

        # Title
        ttk.Label(
            inner,
            text=f"Roll not found in master table!\nBatch No.: {charge}",
            font=("Arial", 22, "bold"),
            foreground="red",
            justify=tk.CENTER,
        ).pack(pady=(0, 16))

        # Pre-filled / read-only info
        info_frame = ttk.LabelFrame(inner, text="Data from QR Code (read-only)", padding="10")
        info_frame.pack(fill=tk.X, pady=(0, 12))
        info_frame.columnconfigure(1, weight=1)

        pre_filled = [
            ("Batch No.:", str(self._qr.get("charge", charge))),
            ("Location (QR):", str(self._qr.get("lort", ""))),
            ("Stage 0 - Length (mm):", str(self._qr.get("lnge0", ""))),
            ("Stage 0 - Width (mm):", str(self._qr.get("brte0", ""))),
            ("Stage 1 - Length (mm):", str(self._qr.get("lnge1", ""))),
            ("Stage 1 - Width (mm):", str(self._qr.get("brte1", ""))),
            ("Stage 2 - Length (mm):", str(self._qr.get("lnge2", ""))),
            ("Stage 2 - Width (mm):", str(self._qr.get("brte2", ""))),
        ]

        for r, (lbl, val) in enumerate(pre_filled):
            ttk.Label(info_frame, text=lbl, font=("Arial", 16, "bold")).grid(
                row=r, column=0, sticky=tk.W, pady=3)
            ttk.Label(info_frame, text=val, font=("Arial", 16)).grid(
                row=r, column=1, sticky=tk.W, padx=(12, 0), pady=3)

        # Manual fields
        manual_frame = ttk.LabelFrame(inner, text="Manual Input", padding="10")
        manual_frame.pack(fill=tk.X, pady=(0, 12))

        self.material_var = tk.StringVar()
        self.kurztext_var = tk.StringVar()
        self.fach_var = tk.StringVar()
        self.brte_meas_var = tk.StringVar()
        self.remarks_var = tk.StringVar()

        fields = [
            ("Material No. *:", self.material_var, True),
            ("Description:", self.kurztext_var, False),
            ("Shelf Location *:", self.fach_var, True),
            ("Measured Width (mm) *:", self.brte_meas_var, True),
            ("Remarks:", self.remarks_var, False),
        ]

        self._entries = {}
        for i, (label, var, _required) in enumerate(fields):
            field_frame = ttk.Frame(manual_frame)
            field_frame.pack(fill=tk.X, pady=6)
            ttk.Label(field_frame, text=label, font=("Arial", 18)).pack(anchor=tk.W)
            entry = ttk.Entry(field_frame, textvariable=var, width=50, font=("Arial", 20))
            entry.pack(fill=tk.X)
            self._entries[label] = entry
            if i == 0:
                entry.focus_set()

        # Buttons
        btn_frame = ttk.Frame(inner)
        btn_frame.pack(pady=(16, 8))
        ttk.Button(btn_frame, text="Save", command=self._save, width=14).pack(
            side=tk.LEFT, padx=(0, 16))
        ttk.Button(btn_frame, text="Cancel", command=self._cancel, width=14).pack(
            side=tk.LEFT)

        self.dialog.bind("<Escape>", lambda e: self._cancel())

    def _save(self):
        material = self.material_var.get().strip()
        fach = self.fach_var.get().strip()
        brte_meas = self.brte_meas_var.get().strip()

        if not material:
            messagebox.showerror("Error", "Material is a required field.", parent=self.dialog)
            return
        if not fach:
            messagebox.showerror("Error", "Shelf Location is a required field.", parent=self.dialog)
            return
        if not brte_meas:
            messagebox.showerror("Error", "Measured Width (mm) is a required field.", parent=self.dialog)
            return

        self.result = {
            "charge": str(self._qr.get("charge", "")),
            "lort_master": "",
            "lort_qr": str(self._qr.get("lort", "")),
            "material": material,
            "kurztext": self.kurztext_var.get().strip(),
            "werk": "",
            "lnge0": self._qr.get("lnge0", 0),
            "brte0": self._qr.get("brte0", 0),
            "lnge1": self._qr.get("lnge1", 0),
            "brte1": self._qr.get("brte1", 0),
            "lnge2": self._qr.get("lnge2", 0),
            "brte2": self._qr.get("brte2", 0),
            "flache": "",
            "frei_verw": "",
            "fach": fach,
            "brte_meas": brte_meas,
            "remarks": self.remarks_var.get().strip(),
            "status": "not_found",
        }
        self.dialog.destroy()

    def _cancel(self):
        self.result = None
        self.dialog.destroy()


# ---------------------------------------------------------------------------
# NotFoundDialogZert
# ---------------------------------------------------------------------------

class NotFoundDialogZert:
    """Dialog for Zert warehouse: batch not found in master table.
    User supplies Material No. (mandatory), Quantity (mandatory), Remarks (optional).
    """

    def __init__(self, parent, charge):
        self.result = None

        self.dialog = tk.Toplevel(parent)
        self.dialog.title("Batch Not Found - Manual Entry")
        self.dialog.geometry("720x520")
        self.dialog.resizable(True, True)
        self.dialog.transient(parent)
        self.dialog.grab_set()
        self.dialog.geometry(
            "+%d+%d" % (parent.winfo_rootx() + 60, parent.winfo_rooty() + 60)
        )
        self._charge = charge
        self._build_widgets()
        self.dialog.wait_window()

    def _build_widgets(self):
        frame = ttk.Frame(self.dialog, padding="20")
        frame.pack(fill=tk.BOTH, expand=True)

        ttk.Label(
            frame,
            text=f"Batch not found!\nBatch No.: {self._charge}",
            font=("Arial", 20, "bold"),
            foreground="red",
            justify=tk.CENTER,
        ).pack(pady=(0, 16))

        info_frame = ttk.LabelFrame(frame, text="QR Data (read-only)", padding="8")
        info_frame.pack(fill=tk.X, pady=(0, 12))
        ttk.Label(info_frame, text="Batch No.:", font=("Arial", 16, "bold")).grid(row=0, column=0, sticky=tk.W)
        ttk.Label(info_frame, text=self._charge, font=("Arial", 16)).grid(row=0, column=1, sticky=tk.W, padx=(8, 0))

        manual_frame = ttk.LabelFrame(frame, text="Manual Input", padding="8")
        manual_frame.pack(fill=tk.X, pady=(0, 16))

        self.material_var = tk.StringVar()
        self.menge_var = tk.StringVar()
        self.remarks_var = tk.StringVar()

        ttk.Label(manual_frame, text="Material No. *:", font=("Arial", 18)).grid(row=0, column=0, sticky=tk.W, pady=6)
        mat_entry = ttk.Entry(manual_frame, textvariable=self.material_var, width=28, font=("Arial", 20))
        mat_entry.grid(row=0, column=1, sticky=tk.W, padx=(12, 0), pady=6)
        mat_entry.focus_set()

        ttk.Label(manual_frame, text="Recorded Quantity *:", font=("Arial", 18)).grid(row=1, column=0, sticky=tk.W, pady=6)
        menge_entry = ttk.Entry(manual_frame, textvariable=self.menge_var, width=14, font=("Arial", 20))
        menge_entry.grid(row=1, column=1, sticky=tk.W, padx=(12, 0), pady=6)

        ttk.Label(manual_frame, text="Remarks:", font=("Arial", 18)).grid(row=2, column=0, sticky=tk.W, pady=6)
        ttk.Entry(manual_frame, textvariable=self.remarks_var, width=28, font=("Arial", 20)).grid(row=2, column=1, sticky=tk.W, padx=(12, 0), pady=6)

        btn_frame = ttk.Frame(frame)
        btn_frame.pack(pady=(12, 0))
        ttk.Button(btn_frame, text="Save", command=self._save, width=14).pack(side=tk.LEFT, padx=(0, 16))
        ttk.Button(btn_frame, text="Cancel", command=self._cancel, width=14).pack(side=tk.LEFT)

        self.dialog.bind("<Escape>", lambda e: self._cancel())

    def _save(self):
        material = self.material_var.get().strip()
        menge = self.menge_var.get().strip()
        if not material:
            messagebox.showerror("Error", "Material No. is a required field.", parent=self.dialog)
            return
        if not menge:
            messagebox.showerror("Error", "Recorded Quantity is a required field.", parent=self.dialog)
            return
        self.result = {
            "charge": self._charge,
            "material": material,
            "kurztext": "",
            "mart": "",
            "werk": "",
            "lort": "",
            "bme": "",
            "frei_verw": "",
            "laenge_mm": "",
            "breite_mm": "",
            "adv": "",
            "menge": menge,
            "remarks": self.remarks_var.get().strip(),
            "status": "not_found",
        }
        self.dialog.destroy()

    def _cancel(self):
        self.result = None
        self.dialog.destroy()


# ---------------------------------------------------------------------------
# NotFoundDialogKMAT
# ---------------------------------------------------------------------------

class NotFoundDialogKMAT:
    """Dialog for KMAT: scanned Order not found in master table."""

    def __init__(self, parent, order):
        self.result = None
        self.dialog = tk.Toplevel(parent)
        self.dialog.title("Not Found - Manual Entry")
        self.dialog.geometry("480x300")
        self.dialog.resizable(False, False)
        self.dialog.transient(parent)
        self.dialog.grab_set()
        self.dialog.geometry(
            "+%d+%d" % (parent.winfo_rootx() + 60, parent.winfo_rooty() + 60)
        )
        self._order = order
        self._build_widgets()
        self.dialog.wait_window()

    def _build_widgets(self):
        frame = ttk.Frame(self.dialog, padding="20")
        frame.pack(fill=tk.BOTH, expand=True)

        ttk.Label(
            frame,
            text=f"Order not found!\nOrder: {self._order}",
            font=("Arial", 12, "bold"),
            foreground="red",
            justify=tk.CENTER,
        ).pack(pady=(0, 16))

        manual_frame = ttk.LabelFrame(frame, text="Manual Input", padding="8")
        manual_frame.pack(fill=tk.X, pady=(0, 12))

        self.material_var = tk.StringVar()
        self.menge_var = tk.StringVar()
        self.remarks_var = tk.StringVar()

        ttk.Label(manual_frame, text="Material No. *:", font=("Arial", 9)).grid(row=0, column=0, sticky=tk.W, pady=4)
        mat_entry = ttk.Entry(manual_frame, textvariable=self.material_var, width=30, font=("Arial", 10))
        mat_entry.grid(row=0, column=1, sticky=tk.W, padx=(8, 0), pady=4)
        mat_entry.focus_set()

        ttk.Label(manual_frame, text="Recorded Quantity *:", font=("Arial", 9)).grid(row=1, column=0, sticky=tk.W, pady=4)
        ttk.Entry(manual_frame, textvariable=self.menge_var, width=15, font=("Arial", 10)).grid(row=1, column=1, sticky=tk.W, padx=(8, 0), pady=4)

        ttk.Label(manual_frame, text="Remarks:", font=("Arial", 9)).grid(row=2, column=0, sticky=tk.W, pady=4)
        ttk.Entry(manual_frame, textvariable=self.remarks_var, width=30, font=("Arial", 10)).grid(row=2, column=1, sticky=tk.W, padx=(8, 0), pady=4)

        btn_frame = ttk.Frame(frame)
        btn_frame.pack(pady=(8, 0))
        ttk.Button(btn_frame, text="Save", command=self._save, width=12).pack(side=tk.LEFT, padx=(0, 12))
        ttk.Button(btn_frame, text="Cancel", command=self._cancel, width=12).pack(side=tk.LEFT)

        self.dialog.bind("<Escape>", lambda e: self._cancel())

    def _save(self):
        material = self.material_var.get().strip()
        menge = self.menge_var.get().strip()
        if not material:
            messagebox.showerror("Error", "Material No. is required.", parent=self.dialog)
            return
        if not menge:
            messagebox.showerror("Error", "Recorded Quantity is required.", parent=self.dialog)
            return
        self.result = {
            "order": self._order,
            "sales_order": "",
            "material": material,
            "kurztext": "",
            "werk": "",
            "bme": "",
            "menge": menge,
            "remarks": self.remarks_var.get().strip(),
            "status": "not_found",
            "_mode": "kmat",
        }
        self.dialog.destroy()

    def _cancel(self):
        self.result = None
        self.dialog.destroy()


# ---------------------------------------------------------------------------
# NotFoundDialogWIP
# ---------------------------------------------------------------------------

class NotFoundDialogWIP:
    """Dialog for WIP: scanned Order not found in master table."""

    def __init__(self, parent, order):
        self.result = None
        self.dialog = tk.Toplevel(parent)
        self.dialog.title("Not Found - Manual Entry")
        self.dialog.geometry("480x300")
        self.dialog.resizable(False, False)
        self.dialog.transient(parent)
        self.dialog.grab_set()
        self.dialog.geometry(
            "+%d+%d" % (parent.winfo_rootx() + 60, parent.winfo_rooty() + 60)
        )
        self._order = order
        self._build_widgets()
        self.dialog.wait_window()

    def _build_widgets(self):
        frame = ttk.Frame(self.dialog, padding="20")
        frame.pack(fill=tk.BOTH, expand=True)

        ttk.Label(
            frame,
            text=f"Order not found!\nOrder: {self._order}",
            font=("Arial", 12, "bold"),
            foreground="red",
            justify=tk.CENTER,
        ).pack(pady=(0, 16))

        manual_frame = ttk.LabelFrame(frame, text="Manual Input", padding="8")
        manual_frame.pack(fill=tk.X, pady=(0, 12))

        self.material_var = tk.StringVar()
        self.menge_var = tk.StringVar()
        self.remarks_var = tk.StringVar()

        ttk.Label(manual_frame, text="Material No. *:", font=("Arial", 9)).grid(row=0, column=0, sticky=tk.W, pady=4)
        mat_entry = ttk.Entry(manual_frame, textvariable=self.material_var, width=30, font=("Arial", 10))
        mat_entry.grid(row=0, column=1, sticky=tk.W, padx=(8, 0), pady=4)
        mat_entry.focus_set()

        ttk.Label(manual_frame, text="Recorded Quantity *:", font=("Arial", 9)).grid(row=1, column=0, sticky=tk.W, pady=4)
        ttk.Entry(manual_frame, textvariable=self.menge_var, width=15, font=("Arial", 10)).grid(row=1, column=1, sticky=tk.W, padx=(8, 0), pady=4)

        ttk.Label(manual_frame, text="Remarks:", font=("Arial", 9)).grid(row=2, column=0, sticky=tk.W, pady=4)
        ttk.Entry(manual_frame, textvariable=self.remarks_var, width=30, font=("Arial", 10)).grid(row=2, column=1, sticky=tk.W, padx=(8, 0), pady=4)

        btn_frame = ttk.Frame(frame)
        btn_frame.pack(pady=(8, 0))
        ttk.Button(btn_frame, text="Save", command=self._save, width=12).pack(side=tk.LEFT, padx=(0, 12))
        ttk.Button(btn_frame, text="Cancel", command=self._cancel, width=12).pack(side=tk.LEFT)

        self.dialog.bind("<Escape>", lambda e: self._cancel())

    def _save(self):
        material = self.material_var.get().strip()
        menge = self.menge_var.get().strip()
        if not material:
            messagebox.showerror("Error", "Material No. is required.", parent=self.dialog)
            return
        if not menge:
            messagebox.showerror("Error", "Recorded Quantity is required.", parent=self.dialog)
            return
        self.result = {
            "order": self._order,
            "sales_order": "",
            "material": material,
            "kurztext": "",
            "werk": "",
            "bme": "",
            "order_qty": "",
            "basic_finish": "",
            "menge": menge,
            "remarks": self.remarks_var.get().strip(),
            "status": "not_found",
            "_mode": "wip",
        }
        self.dialog.destroy()

    def _cancel(self):
        self.result = None
        self.dialog.destroy()


# ---------------------------------------------------------------------------
# NotFoundDialogGranules
# ---------------------------------------------------------------------------

class NotFoundDialogGranules:
    """Dialog for Granules: batch not found in master table.
    User supplies Material No. (mandatory), Description (optional),
    Counted Qty in KG (mandatory, numeric) and Remarks (optional).
    """

    def __init__(self, parent, charge):
        self.result = None

        self.dialog = tk.Toplevel(parent)
        self.dialog.title("Batch Not Found - Manual Entry")
        self.dialog.geometry("760x600")
        self.dialog.resizable(True, True)
        self.dialog.transient(parent)
        self.dialog.grab_set()
        self.dialog.geometry(
            "+%d+%d" % (parent.winfo_rootx() + 60, parent.winfo_rooty() + 60)
        )
        self._charge = charge
        self._build_widgets()
        self.dialog.wait_window()

    def _build_widgets(self):
        frame = ttk.Frame(self.dialog, padding="20")
        frame.pack(fill=tk.BOTH, expand=True)

        ttk.Label(
            frame,
            text=f"Batch not found!\nBatch No.: {self._charge}",
            font=("Arial", 20, "bold"),
            foreground="red",
            justify=tk.CENTER,
        ).pack(pady=(0, 16))

        manual_frame = ttk.LabelFrame(frame, text="Manual Input", padding="8")
        manual_frame.pack(fill=tk.X, pady=(0, 16))

        self.material_var = tk.StringVar()
        self.kurztext_var = tk.StringVar()
        self.menge_var = tk.StringVar()
        self.remarks_var = tk.StringVar()

        fields = [
            ("Material No. *:", self.material_var, 28),
            ("Description:", self.kurztext_var, 28),
            ("Counted Qty (KG) *:", self.menge_var, 14),
            ("Remarks:", self.remarks_var, 28),
        ]
        entries = []
        for r, (label, var, width) in enumerate(fields):
            ttk.Label(manual_frame, text=label, font=("Arial", 18)).grid(
                row=r, column=0, sticky=tk.W, pady=6)
            entry = ttk.Entry(manual_frame, textvariable=var, width=width, font=("Arial", 20))
            entry.grid(row=r, column=1, sticky=tk.W, padx=(12, 0), pady=6)
            entries.append(entry)
        entries[0].focus_set()
        for entry, nxt in zip(entries, entries[1:]):
            entry.bind("<Return>", lambda e, n=nxt: n.focus_set())
        entries[-1].bind("<Return>", lambda e: self._save())

        btn_frame = ttk.Frame(frame)
        btn_frame.pack(pady=(12, 0))
        ttk.Button(btn_frame, text="Save", command=self._save, width=14).pack(side=tk.LEFT, padx=(0, 16))
        ttk.Button(btn_frame, text="Cancel", command=self._cancel, width=14).pack(side=tk.LEFT)

        self.dialog.bind("<Escape>", lambda e: self._cancel())

    def _save(self):
        material = self.material_var.get().strip()
        menge = self.menge_var.get().strip()
        if not material:
            messagebox.showerror("Error", "Material No. is a required field.", parent=self.dialog)
            return
        if not menge:
            messagebox.showerror("Error", "Counted Qty (KG) is a required field.", parent=self.dialog)
            return
        if _to_float(menge) is None:
            messagebox.showerror("Error", "Counted Qty (KG) must be a number.", parent=self.dialog)
            return
        self.result = {
            "charge": self._charge,
            "material": material,
            "kurztext": self.kurztext_var.get().strip(),
            "werk": "",
            "lort": "",
            "bme": "KG",
            "frei_verw": "",
            "menge": menge,
            "remarks": self.remarks_var.get().strip(),
            "status": "not_found",
            "_mode": "granules",
        }
        self.dialog.destroy()

    def _cancel(self):
        self.result = None
        self.dialog.destroy()


# ---------------------------------------------------------------------------
# Main Application Class
# ---------------------------------------------------------------------------

class InventurApp:
    """Forbo inventory application. One warehouse mode per session (see MODES)."""

    # ------------------------------------------------------------------
    # Initialisation
    # ------------------------------------------------------------------

    def __init__(self):
        self.root = tk.Tk()

        # Show warehouse selection before anything else
        self.root.withdraw()
        sel_dlg = WarehouseSelectionDialog(self.root)
        if sel_dlg.result is None:
            self.root.destroy()
            sys.exit(0)
        self.warehouse_mode = sel_dlg.result
        self.root.deiconify()

        # Base paths (EXE-compatible)
        self.base_dir = Path(self.get_base_path())
        self.config_dir = self.base_dir / "config"
        self.config_dir.mkdir(exist_ok=True)

        # Logging (needs config_dir first)
        self.setup_logging()

        # Config (may override paths)
        self.load_config()

        # Resolve arbeitstabelle / export paths from config
        self._resolve_paths()

        # Data
        self.init_data()

        # UI
        self.setup_ui()

        # Check paths & load master table
        self.check_paths_on_startup()

        # Load existing session
        if self.warehouse_mode == "Zert":
            self.load_existing_zert()
        elif self.warehouse_mode == "KMAT":
            self.load_existing_kmat()
        elif self.warehouse_mode == "WIP":
            self.load_existing_wip()
        elif self.warehouse_mode == "Granules":
            self.load_existing_granules()
        else:
            self.load_existing_halb()

        # Shortcuts
        self.bind_shortcuts()

    # ------------------------------------------------------------------
    # Path / Base helpers
    # ------------------------------------------------------------------

    def get_base_path(self):
        """Return base path — works for both .py script and PyInstaller .exe."""
        if getattr(sys, "frozen", False):
            return os.path.dirname(sys.executable)
        return os.path.dirname(os.path.abspath(__file__))

    def _resolve_paths(self):
        """Derive runtime paths from config."""
        # HALB-type paths (HALB, Assembling, Service, Warehouse share the roll workflow)
        self.halb_mode = self.warehouse_mode if self.warehouse_mode in HALB_MODES else "HALB"
        self.master_key, self.export_key = MODE_CONFIG_KEYS[self.halb_mode]

        arb = self.config.get(self.master_key, "")
        self.arbeitstabelle_path = Path(arb) if arb else None

        exp = self.config.get(self.export_key, "")
        if exp:
            self.export_path = Path(exp)
        else:
            self.export_path = self.base_dir

        self.inventur_path = self.export_path / f"Inventory_{self.halb_mode}.xlsx"
        backup_dir = self.export_path / "backups"
        try:
            backup_dir.mkdir(parents=True, exist_ok=True)
        except Exception:
            pass

        # Zert paths
        zert_arb = self.config.get("arbeitstabelle_zert_path", "")
        self.arbeitstabelle_zert_path = Path(zert_arb) if zert_arb else None

        zert_exp = self.config.get("export_zert_path", "")
        self.export_zert_path = Path(zert_exp) if zert_exp else self.base_dir
        self.inventur_zert_path = self.export_zert_path / "Inventory_Zert.xlsx"

        # KMAT paths
        kmat_arb = self.config.get("arbeitstabelle_kmat_path", "")
        self.arbeitstabelle_kmat_path = Path(kmat_arb) if kmat_arb else None

        kmat_exp = self.config.get("export_kmat_path", "")
        self.export_kmat_path = Path(kmat_exp) if kmat_exp else self.base_dir
        self.inventur_kmat_path = self.export_kmat_path / "Inventory_KMAT.xlsx"

        # WIP paths
        wip_arb = self.config.get("arbeitstabelle_wip_path", "")
        self.arbeitstabelle_wip_path = Path(wip_arb) if wip_arb else None
        wip_exp = self.config.get("export_wip_path", "")
        self.export_wip_path = Path(wip_exp) if wip_exp else self.base_dir
        self.inventur_wip_path = self.export_wip_path / "Inventory_WIP.xlsx"

        # Granules paths
        gran_arb = self.config.get("arbeitstabelle_granules_path", "")
        self.arbeitstabelle_granules_path = Path(gran_arb) if gran_arb else None
        gran_exp = self.config.get("export_granules_path", "")
        self.export_granules_path = Path(gran_exp) if gran_exp else self.base_dir
        self.inventur_granules_path = self.export_granules_path / "Inventory_Granules.xlsx"

    # ------------------------------------------------------------------
    # Logging
    # ------------------------------------------------------------------

    def setup_logging(self):
        """Configure logging to config/inventory.log."""
        log_file = self.config_dir / "inventory.log"
        logging.basicConfig(
            level=logging.INFO,
            format="%(asctime)s - %(levelname)s - %(message)s",
            handlers=[
                logging.FileHandler(log_file, encoding="utf-8"),
                logging.StreamHandler(),
            ],
        )
        self.logger = logging.getLogger(__name__)
        self.logger.info(
            f"Inventory application v{APP_VERSION} started "
            f"(mode: {getattr(self, 'warehouse_mode', 'unknown')})")

    # ------------------------------------------------------------------
    # Config
    # ------------------------------------------------------------------

    def load_config(self):
        """Load config from config/settings.json."""
        config_path = self.config_dir / "settings.json"
        default_config = {
            "auto_save": True,
            "vollbild": True,
            "show_switch_warehouse": False,
        }
        for master_key, export_key in MODE_CONFIG_KEYS.values():
            default_config[master_key] = ""
            default_config[export_key] = ""

        # Migrate settings from the former file name (settings_sk.json)
        legacy_path = self.config_dir / "settings_sk.json"
        if not config_path.exists() and legacy_path.exists():
            try:
                shutil.copy2(legacy_path, config_path)
                self.logger.info(f"Settings migrated from {legacy_path.name} to {config_path.name}")
            except Exception as e:
                self.logger.error(f"Error migrating settings: {e}")

        try:
            if config_path.exists():
                with open(config_path, "r", encoding="utf-8-sig") as f:
                    self.config = json.load(f)
                # Merge any missing keys
                for k, v in default_config.items():
                    self.config.setdefault(k, v)
            else:
                self.config = default_config
                self.save_config()
        except Exception as e:
            self.config = default_config
            self.logger.error(f"Error loading config: {e}")

    def save_config(self):
        """Save config to config/settings.json."""
        try:
            config_path = self.config_dir / "settings.json"
            with open(config_path, "w", encoding="utf-8") as f:
                json.dump(self.config, f, indent=2, ensure_ascii=False)
        except Exception as e:
            self.logger.error(f"Error saving config: {e}")

    # ------------------------------------------------------------------
    # Data initialisation
    # ------------------------------------------------------------------

    def init_data(self):
        """Initialise data structures."""
        self.df_rollen = None
        self.inventur_data = []       # found rolls (HALB)
        self.not_found_data = []      # not-found rolls (HALB)
        self.current_scan = None
        self.current_qr_data = None
        self.undo_stack = []

        # Zert data
        self.df_zert = None
        self.inventur_data_zert = []
        self.not_found_data_zert = []

        # KMAT data
        self.df_kmat = None
        self.inventur_data_kmat = []
        self.not_found_data_kmat = []

        # WIP data
        self.inventur_data_wip = []
        self.not_found_data_wip = []
        self.df_wip = None

        # Granules data
        self.df_granules = None
        self.inventur_data_granules = []
        self.not_found_data_granules = []

    # ------------------------------------------------------------------
    # QR parsing
    # ------------------------------------------------------------------

    def parse_qr_code(self, raw):
        """Parse semicolon-delimited QR code.

        Format: Lort;Charge;Lnge0;Brte0;Lnge1;Brte1;Lnge2;Brte2
        Fallback: treat whole string as Charge.

        Returns dict with keys:
            lort, charge, lnge0, brte0, lnge1, brte1, lnge2, brte2
        """
        raw = raw.strip()
        if ";" in raw:
            parts = raw.split(";")

            def _int(val, default=0):
                try:
                    return int(val.strip())
                except (ValueError, AttributeError):
                    return default

            return {
                "lort": parts[0].strip() if len(parts) > 0 else "",
                "charge": parts[1].strip() if len(parts) > 1 else raw,
                "lnge0": _int(parts[2]) if len(parts) > 2 else 0,
                "brte0": _int(parts[3]) if len(parts) > 3 else 0,
                "lnge1": _int(parts[4]) if len(parts) > 4 else 0,
                "brte1": _int(parts[5]) if len(parts) > 5 else 0,
                "lnge2": _int(parts[6]) if len(parts) > 6 else 0,
                "brte2": _int(parts[7]) if len(parts) > 7 else 0,
            }
        else:
            return {
                "lort": "",
                "charge": raw,
                "lnge0": 0,
                "brte0": 0,
                "lnge1": 0,
                "brte1": 0,
                "lnge2": 0,
                "brte2": 0,
            }

    # ------------------------------------------------------------------
    # Master table (HALB)
    # ------------------------------------------------------------------

    def load_arbeitstabelle(self):
        """Load HALB-type master table from the configured path (single sheet)."""
        if self.arbeitstabelle_path is None or not self.arbeitstabelle_path.exists():
            self.df_rollen = None
            self.logger.warning("Master table not loaded (path not set or file missing)")
            return

        try:
            self.df_rollen = pd.read_excel(
                self.arbeitstabelle_path,
                dtype={"Batch": str},
            )
            if "Batch" in self.df_rollen.columns:
                self.df_rollen["Batch"] = self.df_rollen["Batch"].astype(str)

            count = len(self.df_rollen)
            self.logger.info(f"Master table loaded: {count} rows from {self.arbeitstabelle_path}")
            self.status_var.set(f"Master table loaded: {count} rolls")

            # Refresh header info label
            self._update_header_info()

        except Exception as e:
            messagebox.showerror("Error", f"Error loading master table:\n{e}")
            self.logger.error(f"Error loading master table: {e}")
            self.df_rollen = None

    # ------------------------------------------------------------------
    # Master table (Zert)
    # ------------------------------------------------------------------

    def load_arbeitstabelle_zert(self):
        """Load Zert master table from the configured path."""
        if self.arbeitstabelle_zert_path is None or not self.arbeitstabelle_zert_path.exists():
            self.df_zert = None
            self.logger.warning("Zert master table not loaded")
            return
        try:
            self.df_zert = pd.read_excel(self.arbeitstabelle_zert_path, dtype={"Batch": str})
            if "Batch" in self.df_zert.columns:
                self.df_zert["Batch"] = self.df_zert["Batch"].astype(str)
            count = len(self.df_zert)
            self.logger.info(f"Zert master table loaded: {count} rows")
            self.status_var.set(f"Zert master table loaded: {count} entries")
            self._update_header_info()
        except Exception as e:
            messagebox.showerror("Error", f"Error loading Zert master table:\n{e}")
            self.df_zert = None

    # ------------------------------------------------------------------
    # Master table (KMAT)
    # ------------------------------------------------------------------

    def load_arbeitstabelle_kmat(self):
        """Load KMAT master table from the configured path."""
        if self.arbeitstabelle_kmat_path is None or not self.arbeitstabelle_kmat_path.exists():
            self.df_kmat = None
            self.logger.warning("KMAT master table not loaded")
            return
        try:
            self.df_kmat = pd.read_excel(self.arbeitstabelle_kmat_path, dtype=str)

            required = {"Sales Order", "Order", "Material Number"}
            missing = required - set(self.df_kmat.columns)
            if missing:
                raise ValueError(
                    f"KMAT master table is missing required column(s): {', '.join(sorted(missing))}. "
                )

            self.df_kmat["Sales Order"] = self.df_kmat["Sales Order"].apply(_norm_num)
            self.df_kmat["Order"] = self.df_kmat["Order"].apply(_norm_num)

            count = len(self.df_kmat)
            self.logger.info(f"KMAT master table loaded: {count} rows")
            self.status_var.set(f"KMAT master table loaded: {count} entries")
            self._update_header_info()
        except Exception as e:
            messagebox.showerror("Error", f"Error loading KMAT master table:\n{e}")
            self.df_kmat = None

    # ------------------------------------------------------------------
    # Master table (WIP)
    # ------------------------------------------------------------------

    def load_arbeitstabelle_wip(self):
        """Load WIP master table from the configured path."""
        if self.arbeitstabelle_wip_path is None or not self.arbeitstabelle_wip_path.exists():
            self.df_wip = None
            self.logger.warning("WIP master table not loaded")
            return
        try:
            self.df_wip = pd.read_excel(
                self.arbeitstabelle_wip_path,
                sheet_name=0,
                dtype=str,
            )

            if "Sales Order" in self.df_wip.columns:
                self.df_wip["Sales Order"] = self.df_wip["Sales Order"].apply(_norm_num)
            if "Order" in self.df_wip.columns:
                self.df_wip["Order"] = self.df_wip["Order"].apply(_norm_num)

            count = len(self.df_wip)
            self.logger.info(f"WIP master table loaded: {count} rows")
            self.status_var.set(f"WIP master table loaded: {count} entries")
            self._update_header_info()
        except Exception as e:
            messagebox.showerror("Error", f"Error loading WIP master table:\n{e}")
            self.df_wip = None

    # ------------------------------------------------------------------
    # Master table (Granules)
    # ------------------------------------------------------------------

    def load_arbeitstabelle_granules(self):
        """Load Granules master table (German or English headers, see GRANULES_COLUMN_MAP)."""
        if self.arbeitstabelle_granules_path is None or not self.arbeitstabelle_granules_path.exists():
            self.df_granules = None
            self.logger.warning("Granules master table not loaded")
            return
        try:
            df = pd.read_excel(self.arbeitstabelle_granules_path, sheet_name=0, dtype=str)
            df.columns = [str(c).strip() for c in df.columns]
            df = df.rename(columns={c: GRANULES_COLUMN_MAP[c] for c in df.columns if c in GRANULES_COLUMN_MAP})

            if "Batch" not in df.columns:
                raise ValueError(
                    "Granules master table is missing the batch column ('Charge' or 'Batch').")

            df["_batch_key"] = df["Batch"].apply(_batch_key)
            self.df_granules = df

            count = len(self.df_granules)
            self.logger.info(f"Granules master table loaded: {count} rows")
            self.status_var.set(f"Granules master table loaded: {count} entries")
            self._update_header_info()
        except Exception as e:
            messagebox.showerror("Error", f"Error loading Granules master table:\n{e}")
            self.df_granules = None

    # ------------------------------------------------------------------
    # Charge lookup (HALB)
    # ------------------------------------------------------------------

    def suche_charge(self, charge):
        """Look up a charge in df_rollen. Returns row dict or None."""
        if self.df_rollen is None:
            return None
        matches = self.df_rollen[self.df_rollen["Batch"] == str(charge)]
        if not matches.empty:
            return matches.iloc[0].to_dict()
        return None

    # ------------------------------------------------------------------
    # Charge lookup (Zert)
    # ------------------------------------------------------------------

    def suche_charge_zert(self, charge):
        """Look up a charge in df_zert. Returns row dict or None."""
        if self.df_zert is None:
            return None
        matches = self.df_zert[self.df_zert["Batch"] == str(charge)]
        if not matches.empty:
            return matches.iloc[0].to_dict()
        return None

    # ------------------------------------------------------------------
    # Kauf+POS lookup (KMAT)
    # ------------------------------------------------------------------

    def suche_kmat(self, order):
        """Look up by Order in df_kmat. Returns row dict or None."""
        if self.df_kmat is None:
            return None
        order_norm = _norm_num_lookup(order)
        matches = self.df_kmat[self.df_kmat["Order"] == order_norm]
        if not matches.empty:
            return matches.iloc[0].to_dict()
        return None

    def suche_wip(self, order):
        """Look up by Order in df_wip. Returns row dict or None."""
        if self.df_wip is None:
            return None
        order_norm = _norm_num_lookup(order)
        matches = self.df_wip[self.df_wip["Order"] == order_norm]
        if not matches.empty:
            return matches.iloc[0].to_dict()
        return None

    def suche_granules(self, charge):
        """Look up a batch in df_granules (leading zeros ignored). Returns row dict or None."""
        if self.df_granules is None:
            return None
        key = _batch_key(charge)
        if not key:
            return None
        matches = self.df_granules[self.df_granules["_batch_key"] == key]
        if not matches.empty:
            return matches.iloc[0].to_dict()
        return None

    # ------------------------------------------------------------------
    # Startup path check
    # ------------------------------------------------------------------

    def check_paths_on_startup(self):
        """Show a warning / prompt to locate file if master path is not set."""
        if self.warehouse_mode == "Zert":
            arb_missing = (
                self.arbeitstabelle_zert_path is None
                or not self.arbeitstabelle_zert_path.exists()
            )

            if arb_missing:
                answer = messagebox.askyesno(
                    "Zert Master Table Not Found",
                    "The Zert master table has not been configured or could not be found.\n\n"
                    "Search now?",
                )
                if answer:
                    path = filedialog.askopenfilename(
                        title="Select Zert Master Table",
                        filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
                    )
                    if path:
                        self.config["arbeitstabelle_zert_path"] = path
                        self.save_config()
                        self._resolve_paths()
                        self.load_arbeitstabelle_zert()
                    else:
                        self._disable_scan("No Zert master table selected. Scanning disabled.")
                else:
                    self._disable_scan("Zert master table not configured. Scanning disabled.")
            else:
                self.load_arbeitstabelle_zert()

            # Default export zert path
            if not self.config.get("export_zert_path", ""):
                self.config["export_zert_path"] = str(self.base_dir)
                self.save_config()
                self._resolve_paths()
        elif self.warehouse_mode == "KMAT":
            arb_missing = (
                self.arbeitstabelle_kmat_path is None
                or not self.arbeitstabelle_kmat_path.exists()
            )
            if arb_missing:
                answer = messagebox.askyesno(
                    "KMAT Master Table Not Found",
                    "The KMAT master table has not been configured or could not be found.\n\nSearch now?",
                )
                if answer:
                    path = filedialog.askopenfilename(
                        title="Select KMAT Master Table",
                        filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
                    )
                    if path:
                        self.config["arbeitstabelle_kmat_path"] = path
                        self.save_config()
                        self._resolve_paths()
                        self.load_arbeitstabelle_kmat()
                    else:
                        self._disable_scan("No KMAT master table selected. Scanning disabled.")
                else:
                    self._disable_scan("KMAT master table not configured. Scanning disabled.")
            else:
                self.load_arbeitstabelle_kmat()

            if not self.config.get("export_kmat_path", ""):
                self.config["export_kmat_path"] = str(self.base_dir)
                self.save_config()
                self._resolve_paths()
        elif self.warehouse_mode == "WIP":
            arb_missing = (
                self.arbeitstabelle_wip_path is None
                or not self.arbeitstabelle_wip_path.exists()
            )
            if arb_missing:
                answer = messagebox.askyesno(
                    "WIP Master Table Not Found",
                    "The WIP master table has not been configured or could not be found.\n\nSearch now?",
                )
                if answer:
                    path = filedialog.askopenfilename(
                        title="Select WIP Master Table",
                        filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
                    )
                    if path:
                        self.config["arbeitstabelle_wip_path"] = path
                        self.save_config()
                        self._resolve_paths()
                        self.load_arbeitstabelle_wip()
                    else:
                        self._disable_scan("No WIP master table selected. Scanning disabled.")
                else:
                    self._disable_scan("WIP master table not configured. Scanning disabled.")
            else:
                self.load_arbeitstabelle_wip()

            if not self.config.get("export_wip_path", ""):
                self.config["export_wip_path"] = str(self.base_dir)
                self.save_config()
                self._resolve_paths()
        elif self.warehouse_mode == "Granules":
            arb_missing = (
                self.arbeitstabelle_granules_path is None
                or not self.arbeitstabelle_granules_path.exists()
            )
            if arb_missing:
                answer = messagebox.askyesno(
                    "Granules Master Table Not Found",
                    "The Granules master table has not been configured or could not be found.\n\nSearch now?",
                )
                if answer:
                    path = filedialog.askopenfilename(
                        title="Select Granules Master Table",
                        filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
                    )
                    if path:
                        self.config["arbeitstabelle_granules_path"] = path
                        self.save_config()
                        self._resolve_paths()
                        self.load_arbeitstabelle_granules()
                    else:
                        self._disable_scan("No Granules master table selected. Scanning disabled.")
                else:
                    self._disable_scan("Granules master table not configured. Scanning disabled.")
            else:
                self.load_arbeitstabelle_granules()

            if not self.config.get("export_granules_path", ""):
                self.config["export_granules_path"] = str(self.base_dir)
                self.save_config()
                self._resolve_paths()
        else:
            mode_label = MODES[self.halb_mode][0]
            arb_missing = (
                self.arbeitstabelle_path is None
                or not self.arbeitstabelle_path.exists()
            )

            if arb_missing:
                answer = messagebox.askyesno(
                    f"{mode_label} Master Table Not Found",
                    f"The {mode_label} master table file has not been configured or could not be found.\n\n"
                    "Would you like to locate it now?",
                )
                if answer:
                    path = filedialog.askopenfilename(
                        title=f"Select {mode_label} Master Table file",
                        filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
                    )
                    if path:
                        self.config[self.master_key] = path
                        self.save_config()
                        self._resolve_paths()
                        self.load_arbeitstabelle()
                    else:
                        self._disable_scan("No master table selected. Scanning disabled.")
                else:
                    self._disable_scan("Master table not configured. Scanning disabled.")
            else:
                self.load_arbeitstabelle()

            # Default export path to base_dir if not set
            if not self.config.get(self.export_key, ""):
                self.config[self.export_key] = str(self.base_dir)
                self.save_config()
                self._resolve_paths()

    def _disable_scan(self, msg):
        self.scan_entry.config(state="disabled")
        self.status_var.set(f"WARNING: {msg}")
        self.logger.warning(msg)

    # ------------------------------------------------------------------
    # Settings dialog
    # ------------------------------------------------------------------

    def open_settings_dialog(self):
        """Open the settings dialog and apply any changes."""
        dlg = SettingsDialog(self.root, self.config, self.warehouse_mode)
        if dlg.result:
            self.config.update(dlg.result)
            self.save_config()
            self._resolve_paths()
            self._apply_switch_button_visibility()
            # Re-enable scan if it was disabled
            self.scan_entry.config(state="normal")
            if self.warehouse_mode == "Zert":
                self.load_arbeitstabelle_zert()
                messagebox.showinfo("Settings Saved", "Settings have been saved.\nZert master table reloaded.")
            elif self.warehouse_mode == "KMAT":
                self.load_arbeitstabelle_kmat()
                messagebox.showinfo("Settings Saved", "Settings have been saved.\nKMAT master table reloaded.")
            elif self.warehouse_mode == "WIP":
                self.load_arbeitstabelle_wip()
                messagebox.showinfo("Settings Saved", "Settings have been saved.\nWIP master table reloaded.")
            elif self.warehouse_mode == "Granules":
                self.load_arbeitstabelle_granules()
                messagebox.showinfo("Settings Saved", "Settings have been saved.\nGranules master table reloaded.")
            else:
                self.load_arbeitstabelle()
                messagebox.showinfo("Settings Saved", "Settings have been saved.\nMaster table reloaded.")

    def _apply_switch_button_visibility(self):
        if self.config.get("show_switch_warehouse", False):
            self.switch_btn.grid()
        else:
            self.switch_btn.grid_remove()

    def switch_warehouse(self):
        """Restart the app so the warehouse selection dialog appears again."""
        if messagebox.askyesno("Switch Warehouse", "Switch to a different warehouse?\nUnsaved data will be lost."):
            self.root.destroy()
            os.execv(sys.executable, [sys.executable] + sys.argv)

    # ------------------------------------------------------------------
    # UI construction
    # ------------------------------------------------------------------

    def setup_ui(self):
        """Build the main UI."""
        style = ttk.Style()
        style.configure("BigScan.TButton", font=("Arial", 20))
        style.configure("BigSave.TButton", font=("Arial", 20))
        self.root.title(f"INVENTORY Forbo - {MODES[self.warehouse_mode][0]}  (v{APP_VERSION})")
        self.root.geometry("1280x820")
        self.root.configure(bg="#f0f0f0")

        try:
            icon_path = Path(__file__).parent / "icon.ico"
            self.root.iconbitmap(str(icon_path))
        except Exception:
            pass

        self.root.state("zoomed")

        self.main_frame = ttk.Frame(self.root, padding="10")
        self.main_frame.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))

        self.root.columnconfigure(0, weight=1)
        self.root.rowconfigure(0, weight=1)
        self.main_frame.columnconfigure(0, weight=1)

        self.create_header()
        self.create_scan_section()
        self.create_current_scan_section()
        self.create_list_section()
        self.create_button_section()
        self.create_status_bar()

        self.scan_entry.focus_set()

    # --- Header ---

    def create_header(self):
        header_frame = ttk.Frame(self.main_frame)
        header_frame.grid(row=0, column=0, sticky=(tk.W, tk.E), pady=(0, 16))
        header_frame.columnconfigure(1, weight=1)

        mode_label, fg_color = MODES[self.warehouse_mode]
        header_text = f"INVENTORY Forbo - {mode_label}"

        ttk.Label(header_frame, text="[F]", font=("Arial", 20, "bold"),
                  foreground=fg_color).grid(row=0, column=0, padx=(0, 16))

        ttk.Label(
            header_frame,
            text=header_text,
            font=("Arial", 18, "bold"),
            foreground=fg_color,
        ).grid(row=0, column=1, sticky=tk.W)

        self.header_info_var = tk.StringVar(value="No master table loaded")
        ttk.Label(header_frame, textvariable=self.header_info_var,
                  font=("Arial", 10)).grid(row=0, column=2, padx=(20, 0))

        ttk.Label(header_frame, text=f"v{APP_VERSION}", font=("Arial", 10),
                  foreground="gray").grid(row=0, column=3, padx=(20, 0))

    def _update_header_info(self):
        if self.warehouse_mode == "Zert":
            if self.df_zert is not None:
                self.header_info_var.set(f"DB: {len(self.df_zert)} entries in Zert master table")
            else:
                self.header_info_var.set("No Zert master table loaded")
        elif self.warehouse_mode == "KMAT":
            if self.df_kmat is not None:
                self.header_info_var.set(f"DB: {len(self.df_kmat)} entries in KMAT master table")
            else:
                self.header_info_var.set("No KMAT master table loaded")
        elif self.warehouse_mode == "WIP":
            if self.df_wip is not None:
                self.header_info_var.set(f"DB: {len(self.df_wip)} entries in WIP master table")
            else:
                self.header_info_var.set("No WIP master table loaded")
        elif self.warehouse_mode == "Granules":
            if self.df_granules is not None:
                self.header_info_var.set(f"DB: {len(self.df_granules)} batches in Granules master table")
            else:
                self.header_info_var.set("No Granules master table loaded")
        else:
            if self.df_rollen is not None:
                self.header_info_var.set(f"DB: {len(self.df_rollen)} rolls in master table")
            else:
                self.header_info_var.set("No master table loaded")

    # --- Scan section ---

    def create_scan_section(self):
        scan_frame = ttk.LabelFrame(self.main_frame, text="BARCODE / QR SCAN", padding="10")
        scan_frame.grid(row=1, column=0, sticky=(tk.W, tk.E), pady=(0, 10))
        scan_frame.columnconfigure(0, weight=1)

        self.scan_var = tk.StringVar()
        self.scan_entry = ttk.Entry(scan_frame, textvariable=self.scan_var,
                                    font=("Arial", 28), width=50)
        self.scan_entry.grid(row=0, column=0, sticky=(tk.W, tk.E), padx=(0, 10))

        ttk.Button(scan_frame, text="Scan", command=self.process_scan, style="BigScan.TButton").grid(row=0, column=1, padx=(0, 0))

        self.scan_entry.bind("<Return>", lambda e: self.process_scan())

        ttk.Label(
            scan_frame,
            text="Scan field is always focused. Scan a QR code or type + ENTER.",
            font=("Arial", 9),
            foreground="gray",
        ).grid(row=1, column=0, columnspan=2, pady=(5, 0))

    # --- Current scan section ---

    def create_current_scan_section(self):
        self.current_frame = ttk.LabelFrame(
            self.main_frame, text="CURRENT SCAN", padding="10")
        self.current_frame.grid(row=2, column=0, sticky=(tk.W, tk.E), pady=(0, 10))
        self.current_frame.columnconfigure(1, weight=1)
        self.current_frame.columnconfigure(3, weight=1)

        self._create_display_labels()
        self._create_input_panel()
        self._create_zert_info_panel()
        self._create_zert_input_panel()
        self._create_kmat_info_panel()
        self._create_kmat_input_panel()
        self._create_wip_info_panel()
        self._create_wip_input_panel()
        self._create_granules_panel()

        self.current_frame.grid_remove()

    def _create_display_labels(self):
        """Create read-only display labels for HALB scan data."""
        # Row 0: Batch No. | Material
        ttk.Label(self.current_frame, text="Batch No.:", font=("Arial", 11, "bold")).grid(
            row=0, column=0, sticky=tk.W, pady=2)
        self.lbl_charge = ttk.Label(self.current_frame, text="", font=("Arial", 11))
        self.lbl_charge.grid(row=0, column=1, sticky=tk.W, padx=(8, 20), pady=2)

        ttk.Label(self.current_frame, text="Material:", font=("Arial", 11, "bold")).grid(
            row=0, column=2, sticky=tk.W, pady=2)
        self.lbl_material = ttk.Label(self.current_frame, text="", font=("Arial", 11))
        self.lbl_material.grid(row=0, column=3, sticky=tk.W, padx=(8, 0), pady=2)

        # Row 1: Description
        ttk.Label(self.current_frame, text="Description:", font=("Arial", 11, "bold")).grid(
            row=1, column=0, sticky=tk.W, pady=2)
        self.lbl_kurztext = ttk.Label(self.current_frame, text="", font=("Arial", 11))
        self.lbl_kurztext.grid(row=1, column=1, columnspan=3, sticky=tk.W, padx=(8, 0), pady=2)

        # Row 2: Plant | Location (Master)
        ttk.Label(self.current_frame, text="Plant:", font=("Arial", 11, "bold")).grid(
            row=2, column=0, sticky=tk.W, pady=2)
        self.lbl_werk = ttk.Label(self.current_frame, text="", font=("Arial", 11))
        self.lbl_werk.grid(row=2, column=1, sticky=tk.W, padx=(8, 20), pady=2)

        ttk.Label(self.current_frame, text="Location (Master):", font=("Arial", 11, "bold")).grid(
            row=2, column=2, sticky=tk.W, pady=2)
        self.lbl_lort_master = ttk.Label(self.current_frame, text="", font=("Arial", 11))
        self.lbl_lort_master.grid(row=2, column=3, sticky=tk.W, padx=(8, 0), pady=2)

        # Row 3: Location QR | Area
        ttk.Label(self.current_frame, text="Location (QR):", font=("Arial", 11, "bold")).grid(
            row=3, column=0, sticky=tk.W, pady=2)
        self.lbl_lort_qr = ttk.Label(self.current_frame, text="", font=("Arial", 11))
        self.lbl_lort_qr.grid(row=3, column=1, sticky=tk.W, padx=(8, 20), pady=2)

        ttk.Label(self.current_frame, text="Area (m2):", font=("Arial", 11, "bold")).grid(
            row=3, column=2, sticky=tk.W, pady=2)
        self.lbl_flache = ttk.Label(self.current_frame, text="", font=("Arial", 11))
        self.lbl_flache.grid(row=3, column=3, sticky=tk.W, padx=(8, 0), pady=2)

        # Row 4: Free Usable
        ttk.Label(self.current_frame, text="Free Usable:", font=("Arial", 11, "bold")).grid(
            row=4, column=0, sticky=tk.W, pady=2)
        self.lbl_frei = ttk.Label(self.current_frame, text="", font=("Arial", 11))
        self.lbl_frei.grid(row=4, column=1, sticky=tk.W, padx=(8, 0), pady=2)

        # Row 5: Dimensions grid
        self.dim_frame = ttk.LabelFrame(self.current_frame, text="Dimensions (mm)", padding="6")
        self.dim_frame.grid(row=5, column=0, columnspan=4, sticky=(tk.W, tk.E), pady=(8, 4))
        self.dim_frame.columnconfigure(1, weight=1)
        self.dim_frame.columnconfigure(3, weight=1)
        self.dim_frame.columnconfigure(5, weight=1)

        ttk.Label(self.dim_frame, text="Stage 0", font=("Arial", 10, "bold")).grid(
            row=0, column=0, columnspan=2, sticky=tk.W, padx=(0, 16))
        ttk.Label(self.dim_frame, text="Stage 1", font=("Arial", 10, "bold")).grid(
            row=0, column=2, columnspan=2, sticky=tk.W, padx=(0, 16))
        ttk.Label(self.dim_frame, text="Stage 2", font=("Arial", 10, "bold")).grid(
            row=0, column=4, columnspan=2, sticky=tk.W)

        # Stage 0
        ttk.Label(self.dim_frame, text="Length:", font=("Arial", 9)).grid(
            row=1, column=0, sticky=tk.W, padx=(0, 4))
        self.lbl_lnge0 = ttk.Label(self.dim_frame, text="", font=("Arial", 10, "bold"), width=8)
        self.lbl_lnge0.grid(row=1, column=1, sticky=tk.W, padx=(0, 16))

        # Stage 1
        ttk.Label(self.dim_frame, text="Length:", font=("Arial", 9)).grid(
            row=1, column=2, sticky=tk.W, padx=(0, 4))
        self.lbl_lnge1 = ttk.Label(self.dim_frame, text="", font=("Arial", 10, "bold"), width=8)
        self.lbl_lnge1.grid(row=1, column=3, sticky=tk.W, padx=(0, 16))

        # Stage 2
        ttk.Label(self.dim_frame, text="Length:", font=("Arial", 9)).grid(
            row=1, column=4, sticky=tk.W, padx=(0, 4))
        self.lbl_lnge2 = ttk.Label(self.dim_frame, text="", font=("Arial", 10, "bold"), width=8)
        self.lbl_lnge2.grid(row=1, column=5, sticky=tk.W)

        ttk.Label(self.dim_frame, text="Width:", font=("Arial", 9)).grid(
            row=2, column=0, sticky=tk.W, padx=(0, 4))
        self.lbl_brte0 = ttk.Label(self.dim_frame, text="", font=("Arial", 10, "bold"), width=8)
        self.lbl_brte0.grid(row=2, column=1, sticky=tk.W, padx=(0, 16))

        ttk.Label(self.dim_frame, text="Width:", font=("Arial", 9)).grid(
            row=2, column=2, sticky=tk.W, padx=(0, 4))
        self.lbl_brte1 = ttk.Label(self.dim_frame, text="", font=("Arial", 10, "bold"), width=8)
        self.lbl_brte1.grid(row=2, column=3, sticky=tk.W, padx=(0, 16))

        ttk.Label(self.dim_frame, text="Width:", font=("Arial", 9)).grid(
            row=2, column=4, sticky=tk.W, padx=(0, 4))
        self.lbl_brte2 = ttk.Label(self.dim_frame, text="", font=("Arial", 10, "bold"), width=8)
        self.lbl_brte2.grid(row=2, column=5, sticky=tk.W)

        self.halb_separator = ttk.Separator(self.current_frame, orient="horizontal")
        self.halb_separator.grid(row=6, column=0, columnspan=4, sticky=(tk.W, tk.E), pady=8)

    def _create_input_panel(self):
        """Create user input widgets for HALB (Fach + Measured Width + Remarks)."""
        self.fach_var = tk.StringVar()
        self.brte_meas_var = tk.StringVar()
        self.remarks_var = tk.StringVar()
        self.input_widgets = {}

        self.input_container = ttk.Frame(self.current_frame)
        self.input_container.grid(row=7, column=0, columnspan=4,
                                   sticky=(tk.W, tk.E), pady=4)
        self.input_container.columnconfigure(1, weight=1)

        # Fach (mandatory) — row 0
        lbl_fach = ttk.Label(self.input_container, text="Shelf Location *:",
                              font=("Arial", 22, "bold"))
        lbl_fach.grid(row=0, column=0, sticky=tk.W, pady=5)
        entry_fach = ttk.Entry(self.input_container, textvariable=self.fach_var,
                                font=("Arial", 24), width=20)
        entry_fach.grid(row=0, column=1, sticky=tk.W, padx=(10, 0), pady=5)
        entry_fach.bind("<Return>", lambda e: self.input_widgets["brte_meas_entry"].focus_set())
        self.input_widgets["fach_entry"] = entry_fach

        # Measured Width (mandatory) — row 1
        lbl_brte_meas = ttk.Label(self.input_container, text="Measured Width (mm) *:",
                                   font=("Arial", 22, "bold"))
        lbl_brte_meas.grid(row=1, column=0, sticky=tk.W, pady=5)
        entry_brte_meas = ttk.Entry(self.input_container, textvariable=self.brte_meas_var,
                                     font=("Arial", 24), width=10)
        entry_brte_meas.grid(row=1, column=1, sticky=tk.W, padx=(10, 0), pady=5)
        entry_brte_meas.bind("<Return>", self.save_current_scan)
        self.input_widgets["brte_meas_entry"] = entry_brte_meas

        # Remarks (optional) — row 2
        lbl_remarks = ttk.Label(self.input_container, text="Remarks (optional):",
                                 font=("Arial", 22, "bold"))
        lbl_remarks.grid(row=2, column=0, sticky=tk.W, pady=5)
        entry_remarks = ttk.Entry(self.input_container, textvariable=self.remarks_var,
                                   font=("Arial", 24), width=50)
        entry_remarks.grid(row=2, column=1, sticky=tk.W, padx=(10, 0), pady=5)
        entry_remarks.bind("<Return>", self.save_current_scan)
        self.input_widgets["remarks_entry"] = entry_remarks

        # Save button — row 3
        save_btn = ttk.Button(self.input_container, text="Save",
                               command=self.save_current_scan, style="BigSave.TButton")
        save_btn.grid(row=3, column=1, sticky=tk.W, padx=(10, 0), pady=8)
        self.input_widgets["save_button"] = save_btn

    def _create_zert_info_panel(self):
        """Create info display panel for Zert mode (initially hidden)."""
        self.zert_info_container = ttk.LabelFrame(
            self.current_frame, text="Zert Batch Info", padding="8")
        # Don't grid yet — shown on demand

        self.zert_info_container.columnconfigure(1, weight=1)
        self.zert_info_container.columnconfigure(3, weight=1)

        # Labels for Zert info fields
        fields_left = [
            ("Batch No.:", "zert_lbl_charge"),
            ("Material No.:", "zert_lbl_material"),
            ("Description:", "zert_lbl_kurztext"),
            ("Mat. Type:", "zert_lbl_mart"),
            ("Plant:", "zert_lbl_werk"),
            ("Location:", "zert_lbl_lort"),
        ]
        fields_right = [
            ("UOM:", "zert_lbl_bme"),
            ("Free Usable:", "zert_lbl_frei"),
            ("Length (mm):", "zert_lbl_laenge"),
            ("Width (mm):", "zert_lbl_breite"),
            ("ADV:", "zert_lbl_adv"),
        ]

        for r, (lbl_text, attr) in enumerate(fields_left):
            ttk.Label(self.zert_info_container, text=lbl_text,
                      font=("Arial", 10, "bold")).grid(row=r, column=0, sticky=tk.W, pady=2, padx=(0, 6))
            lbl = ttk.Label(self.zert_info_container, text="", font=("Arial", 10))
            lbl.grid(row=r, column=1, sticky=tk.W, padx=(0, 20), pady=2)
            setattr(self, attr, lbl)

        for r, (lbl_text, attr) in enumerate(fields_right):
            ttk.Label(self.zert_info_container, text=lbl_text,
                      font=("Arial", 10, "bold")).grid(row=r, column=2, sticky=tk.W, pady=2, padx=(0, 6))
            lbl = ttk.Label(self.zert_info_container, text="", font=("Arial", 10))
            lbl.grid(row=r, column=3, sticky=tk.W, pady=2)
            setattr(self, attr, lbl)

    def _create_zert_input_panel(self):
        """Create input panel for Zert mode: only Recorded Quantity."""
        self.zert_input_container = ttk.Frame(self.current_frame)
        # Don't grid yet — shown on demand

        self.menge_var = tk.StringVar()
        self.zert_remarks_var = tk.StringVar()
        self.zert_input_widgets = {}

        self.zert_input_container.columnconfigure(1, weight=1)

        ttk.Label(self.zert_input_container, text="Recorded Quantity *:",
                  font=("Arial", 22, "bold")).grid(row=0, column=0, sticky=tk.W, pady=5)
        entry_menge = ttk.Entry(self.zert_input_container, textvariable=self.menge_var,
                                font=("Arial", 24), width=15)
        entry_menge.grid(row=0, column=1, sticky=tk.W, padx=(10, 0), pady=5)
        entry_menge.bind("<Return>", self.save_current_scan_zert)
        self.zert_input_widgets["menge_entry"] = entry_menge

        self.zert_lbl_bme_input = ttk.Label(self.zert_input_container, text="",
                                             font=("Arial", 20), foreground="gray")
        self.zert_lbl_bme_input.grid(row=0, column=2, sticky=tk.W, padx=(8, 0), pady=5)

        ttk.Label(self.zert_input_container, text="Remarks:",
                  font=("Arial", 22, "bold")).grid(row=1, column=0, sticky=tk.W, pady=5)
        entry_zert_remarks = ttk.Entry(self.zert_input_container, textvariable=self.zert_remarks_var,
                                        font=("Arial", 24), width=40)
        entry_zert_remarks.grid(row=1, column=1, columnspan=2, sticky=tk.W, padx=(10, 0), pady=5)
        entry_zert_remarks.bind("<Return>", self.save_current_scan_zert)
        self.zert_input_widgets["remarks_entry"] = entry_zert_remarks

        save_btn = ttk.Button(self.zert_input_container, text="Save",
                               command=self.save_current_scan_zert, style="BigSave.TButton")
        save_btn.grid(row=2, column=1, sticky=tk.W, padx=(10, 0), pady=8)
        self.zert_input_widgets["save_button"] = save_btn

    def _create_kmat_info_panel(self):
        """Create info display panel for KMAT mode (initially hidden)."""
        self.kmat_info_container = ttk.LabelFrame(
            self.current_frame, text="KMAT Info", padding="8")

        self.kmat_info_container.columnconfigure(1, weight=1)
        self.kmat_info_container.columnconfigure(3, weight=1)

        fields_left = [
            ("Sales Order:", "kmat_lbl_sales_order"),
            ("Order:", "kmat_lbl_order"),
            ("Material No.:", "kmat_lbl_material"),
            ("Description:", "kmat_lbl_kurztext"),
        ]
        fields_right = [
            ("Plant:", "kmat_lbl_werk"),
            ("UOM:", "kmat_lbl_bme"),
        ]

        for r, (lbl_text, attr) in enumerate(fields_left):
            ttk.Label(self.kmat_info_container, text=lbl_text,
                      font=("Arial", 10, "bold")).grid(row=r, column=0, sticky=tk.W, pady=2, padx=(0, 6))
            lbl = ttk.Label(self.kmat_info_container, text="", font=("Arial", 10))
            lbl.grid(row=r, column=1, sticky=tk.W, padx=(0, 20), pady=2)
            setattr(self, attr, lbl)

        for r, (lbl_text, attr) in enumerate(fields_right):
            ttk.Label(self.kmat_info_container, text=lbl_text,
                      font=("Arial", 10, "bold")).grid(row=r, column=2, sticky=tk.W, pady=2, padx=(0, 6))
            lbl = ttk.Label(self.kmat_info_container, text="", font=("Arial", 10))
            lbl.grid(row=r, column=3, sticky=tk.W, pady=2)
            setattr(self, attr, lbl)

    def _create_kmat_input_panel(self):
        """Create input panel for KMAT mode."""
        self.kmat_input_container = ttk.Frame(self.current_frame)

        self.kmat_menge_var = tk.StringVar()
        self.kmat_remarks_var = tk.StringVar()
        self.kmat_input_widgets = {}

        self.kmat_input_container.columnconfigure(1, weight=1)

        ttk.Label(self.kmat_input_container, text="Recorded Quantity *:",
                  font=("Arial", 22, "bold")).grid(row=0, column=0, sticky=tk.W, pady=5)
        entry_menge = ttk.Entry(self.kmat_input_container, textvariable=self.kmat_menge_var,
                                font=("Arial", 24), width=15)
        entry_menge.grid(row=0, column=1, sticky=tk.W, padx=(10, 0), pady=5)
        entry_menge.bind("<Return>", self.save_current_scan_kmat)
        self.kmat_input_widgets["menge_entry"] = entry_menge

        self.kmat_lbl_bme_input = ttk.Label(self.kmat_input_container, text="",
                                             font=("Arial", 20), foreground="gray")
        self.kmat_lbl_bme_input.grid(row=0, column=2, sticky=tk.W, padx=(8, 0), pady=5)

        ttk.Label(self.kmat_input_container, text="Remarks:",
                  font=("Arial", 22, "bold")).grid(row=1, column=0, sticky=tk.W, pady=5)
        entry_remarks = ttk.Entry(self.kmat_input_container, textvariable=self.kmat_remarks_var,
                                  font=("Arial", 24), width=40)
        entry_remarks.grid(row=1, column=1, columnspan=2, sticky=tk.W, padx=(10, 0), pady=5)
        entry_remarks.bind("<Return>", self.save_current_scan_kmat)
        self.kmat_input_widgets["remarks_entry"] = entry_remarks

        save_btn = ttk.Button(self.kmat_input_container, text="Save",
                               command=self.save_current_scan_kmat, style="BigSave.TButton")
        save_btn.grid(row=2, column=1, sticky=tk.W, padx=(10, 0), pady=8)
        self.kmat_input_widgets["save_button"] = save_btn

    def _create_wip_info_panel(self):
        """Create info display panel for WIP mode (initially hidden)."""
        self.wip_info_container = ttk.Frame(self.current_frame)
        wip_info = ttk.LabelFrame(self.wip_info_container, text="WIP DETAILS", padding="8")
        wip_info.pack(fill=tk.X)

        wip_info.columnconfigure(1, weight=1)
        wip_info.columnconfigure(3, weight=1)

        ttk.Label(wip_info, text="Sales Order:", font=("Arial", 18)).grid(row=0, column=0, sticky=tk.W, pady=3)
        self.wip_lbl_sales_order = ttk.Label(wip_info, text="", font=("Arial", 20, "bold"))
        self.wip_lbl_sales_order.grid(row=0, column=1, sticky=tk.W, padx=(8, 20), pady=3)

        ttk.Label(wip_info, text="Order:", font=("Arial", 18)).grid(row=0, column=2, sticky=tk.W, pady=3)
        self.wip_lbl_order = ttk.Label(wip_info, text="", font=("Arial", 20, "bold"))
        self.wip_lbl_order.grid(row=0, column=3, sticky=tk.W, padx=(8, 0), pady=3)

        ttk.Label(wip_info, text="Material No.:", font=("Arial", 18)).grid(row=1, column=0, sticky=tk.W, pady=3)
        self.wip_lbl_material = ttk.Label(wip_info, text="", font=("Arial", 20))
        self.wip_lbl_material.grid(row=1, column=1, sticky=tk.W, padx=(8, 20), pady=3)

        ttk.Label(wip_info, text="Plant:", font=("Arial", 18)).grid(row=1, column=2, sticky=tk.W, pady=3)
        self.wip_lbl_werk = ttk.Label(wip_info, text="", font=("Arial", 20))
        self.wip_lbl_werk.grid(row=1, column=3, sticky=tk.W, padx=(8, 0), pady=3)

        ttk.Label(wip_info, text="Description:", font=("Arial", 18)).grid(row=2, column=0, sticky=tk.W, pady=3)
        self.wip_lbl_kurztext = ttk.Label(wip_info, text="", font=("Arial", 20))
        self.wip_lbl_kurztext.grid(row=2, column=1, columnspan=3, sticky=tk.W, padx=(8, 0), pady=3)

        ttk.Label(wip_info, text="Order Qty:", font=("Arial", 18)).grid(row=3, column=0, sticky=tk.W, pady=3)
        self.wip_lbl_order_qty = ttk.Label(wip_info, text="", font=("Arial", 20))
        self.wip_lbl_order_qty.grid(row=3, column=1, sticky=tk.W, padx=(8, 20), pady=3)

        ttk.Label(wip_info, text="UOM:", font=("Arial", 18)).grid(row=3, column=2, sticky=tk.W, pady=3)
        self.wip_lbl_bme = ttk.Label(wip_info, text="", font=("Arial", 20))
        self.wip_lbl_bme.grid(row=3, column=3, sticky=tk.W, padx=(8, 0), pady=3)

        ttk.Label(wip_info, text="Basic Finish:", font=("Arial", 18)).grid(row=4, column=0, sticky=tk.W, pady=3)
        self.wip_lbl_basic_finish = ttk.Label(wip_info, text="", font=("Arial", 20))
        self.wip_lbl_basic_finish.grid(row=4, column=1, sticky=tk.W, padx=(8, 20), pady=3)

    def _create_wip_input_panel(self):
        """Create input panel for WIP mode."""
        self.wip_input_container = ttk.Frame(self.current_frame)
        wip_input = ttk.LabelFrame(self.wip_input_container, text="INPUT", padding="8")
        wip_input.pack(fill=tk.X)

        self.wip_menge_var = tk.StringVar()
        self.wip_remarks_var = tk.StringVar()
        self.wip_input_widgets = {}

        ttk.Label(wip_input, text="Recorded Quantity *:", font=("Arial", 20)).grid(row=0, column=0, sticky=tk.W, pady=4)
        menge_e = ttk.Entry(wip_input, textvariable=self.wip_menge_var, width=15, font=("Arial", 24))
        menge_e.grid(row=0, column=1, sticky=tk.W, padx=(8, 8), pady=4)
        self.wip_input_widgets["menge_entry"] = menge_e

        ttk.Label(wip_input, text="pcs", font=("Arial", 20)).grid(row=0, column=2, sticky=tk.W, pady=4)

        ttk.Label(wip_input, text="Remarks:", font=("Arial", 20)).grid(row=1, column=0, sticky=tk.W, pady=4)
        ttk.Entry(wip_input, textvariable=self.wip_remarks_var, width=30, font=("Arial", 20)).grid(row=1, column=1, columnspan=2, sticky=tk.W, padx=(8, 0), pady=4)

        ttk.Button(wip_input, text="SAVE (Enter)", command=self.save_current_scan_wip, width=16).grid(
            row=2, column=0, columnspan=3, pady=(8, 0))

        menge_e.bind("<Return>", self.save_current_scan_wip)

    def _create_granules_panel(self):
        """Create info + input panel for Granules mode (initially hidden).
        Batch, Material No., Description, Quantity — line — Counted Qty (KG), Remarks."""
        self.granules_container = ttk.Frame(self.current_frame)
        self.granules_container.columnconfigure(1, weight=1)

        info_fields = [
            ("Batch No.:", "gran_lbl_charge"),
            ("Material No.:", "gran_lbl_material"),
            ("Description:", "gran_lbl_kurztext"),
            ("Quantity:", "gran_lbl_frei"),
        ]
        for r, (lbl_text, attr) in enumerate(info_fields):
            ttk.Label(self.granules_container, text=lbl_text,
                      font=("Arial", 18)).grid(row=r, column=0, sticky=tk.W, pady=3)
            lbl = ttk.Label(self.granules_container, text="", font=("Arial", 20, "bold"))
            lbl.grid(row=r, column=1, sticky=tk.W, padx=(10, 0), pady=3)
            setattr(self, attr, lbl)

        ttk.Separator(self.granules_container, orient="horizontal").grid(
            row=4, column=0, columnspan=2, sticky=(tk.W, tk.E), pady=10)

        self.gran_menge_var = tk.StringVar()
        self.gran_remarks_var = tk.StringVar()
        self.gran_input_widgets = {}

        ttk.Label(self.granules_container, text="Counted Qty (KG) *:",
                  font=("Arial", 22, "bold")).grid(row=5, column=0, sticky=tk.W, pady=5)
        entry_menge = ttk.Entry(self.granules_container, textvariable=self.gran_menge_var,
                                font=("Arial", 24), width=15)
        entry_menge.grid(row=5, column=1, sticky=tk.W, padx=(10, 0), pady=5)
        entry_menge.bind("<Return>", self.save_current_scan_granules)
        self.gran_input_widgets["menge_entry"] = entry_menge

        ttk.Label(self.granules_container, text="Remarks (optional):",
                  font=("Arial", 22, "bold")).grid(row=6, column=0, sticky=tk.W, pady=5)
        entry_remarks = ttk.Entry(self.granules_container, textvariable=self.gran_remarks_var,
                                  font=("Arial", 24), width=40)
        entry_remarks.grid(row=6, column=1, sticky=tk.W, padx=(10, 0), pady=5)
        entry_remarks.bind("<Return>", self.save_current_scan_granules)
        self.gran_input_widgets["remarks_entry"] = entry_remarks

        save_btn = ttk.Button(self.granules_container, text="Save",
                              command=self.save_current_scan_granules, style="BigSave.TButton")
        save_btn.grid(row=7, column=1, sticky=tk.W, padx=(10, 0), pady=8)
        self.gran_input_widgets["save_button"] = save_btn

    def _show_granules_widgets(self):
        """Show Granules panel, hide all other mode widgets."""
        for widget in self.current_frame.grid_slaves():
            widget.grid_remove()
        self.granules_container.grid(row=0, column=0, columnspan=4,
                                     sticky=(tk.W, tk.E), pady=4)

    def _show_halb_widgets(self):
        """Show HALB-specific widgets, hide all other mode widgets."""
        # Show HALB rows 0-7
        for r in range(8):
            for widget in self.current_frame.grid_slaves(row=r):
                widget.grid()
        self.dim_frame.grid()
        self.halb_separator.grid()
        self.input_container.grid()
        # Hide Zert containers
        self.zert_info_container.grid_remove()
        self.zert_input_container.grid_remove()
        # Hide KMAT containers
        self.kmat_info_container.grid_remove()
        self.kmat_input_container.grid_remove()
        # Hide WIP / Granules containers
        self.wip_info_container.grid_remove()
        self.wip_input_container.grid_remove()
        self.granules_container.grid_remove()

    def _show_zert_widgets(self):
        """Show Zert-specific widgets, hide all other mode widgets."""
        # Hide all HALB-specific widgets (rows 0-7 labels + dim_frame + separator + input_container)
        for widget in self.current_frame.grid_slaves():
            widget.grid_remove()

        # Show Zert containers
        self.zert_info_container.grid(row=0, column=0, columnspan=4,
                                       sticky=(tk.W, tk.E), pady=(0, 8))
        self.zert_input_container.grid(row=1, column=0, columnspan=4,
                                        sticky=(tk.W, tk.E), pady=4)

    def _show_kmat_widgets(self):
        """Show KMAT-specific widgets, hide all other mode widgets."""
        for widget in self.current_frame.grid_slaves():
            widget.grid_remove()
        self.kmat_info_container.grid(row=0, column=0, columnspan=4,
                                       sticky=(tk.W, tk.E), pady=(0, 8))
        self.kmat_input_container.grid(row=1, column=0, columnspan=4,
                                        sticky=(tk.W, tk.E), pady=4)

    def _show_wip_widgets(self):
        """Show WIP-specific widgets, hide all other mode widgets."""
        for widget in self.current_frame.grid_slaves():
            widget.grid_remove()
        self.wip_info_container.grid(row=0, column=0, columnspan=4,
                                     sticky=(tk.W, tk.E), pady=(0, 8))
        self.wip_input_container.grid(row=1, column=0, columnspan=4,
                                      sticky=(tk.W, tk.E), pady=4)

    def _hide_all_scan_widgets(self):
        """Hide all scan detail widgets."""
        self.zert_info_container.grid_remove()
        self.zert_input_container.grid_remove()
        self.kmat_info_container.grid_remove()
        self.kmat_input_container.grid_remove()
        self.wip_info_container.grid_remove()
        self.wip_input_container.grid_remove()
        self.granules_container.grid_remove()
        self.input_container.grid_remove()

    # --- List section ---

    def create_list_section(self):
        list_frame = ttk.LabelFrame(self.main_frame, text="SCANNED ITEMS", padding="10")
        list_frame.grid(row=3, column=0, sticky=(tk.W, tk.E, tk.N, tk.S), pady=(0, 10))
        list_frame.columnconfigure(0, weight=1)
        list_frame.rowconfigure(1, weight=1)
        self.main_frame.rowconfigure(3, weight=1)

        count_frame = ttk.Frame(list_frame)
        count_frame.grid(row=0, column=0, sticky=(tk.W, tk.E), pady=(0, 8))

        self.count_label = ttk.Label(count_frame, text="0 items", font=("Arial", 12, "bold"))
        self.count_label.pack(side=tk.LEFT)

        if self.warehouse_mode == "Zert":
            columns = ("Time", "Batch No.", "Material No.", "Description", "Quantity", "UOM", "Status")
        elif self.warehouse_mode == "KMAT":
            columns = ("Time", "Sales Order", "Order", "Material No.", "Description", "Quantity", "UOM", "Status")
        elif self.warehouse_mode == "WIP":
            columns = ("Time", "Sales Order", "Order", "Material No.", "Description", "Quantity", "UOM", "Status")
        elif self.warehouse_mode == "Granules":
            columns = ("Time", "Batch No.", "Material No.", "Description", "Counted (KG)", "Free Usable", "Status")
        else:
            columns = ("Time", "Batch No.", "Material", "Shelf Location", "Status")

        self.tree = ttk.Treeview(list_frame, columns=columns, show="headings", height=12)

        for col in columns:
            self.tree.heading(col, text=col)

        if self.warehouse_mode == "Zert":
            self.tree.column("Time", width=90, minwidth=70)
            self.tree.column("Batch No.", width=130, minwidth=100)
            self.tree.column("Material No.", width=110, minwidth=80)
            self.tree.column("Description", width=200, minwidth=120)
            self.tree.column("Quantity", width=80, minwidth=60)
            self.tree.column("UOM", width=60, minwidth=50)
            self.tree.column("Status", width=100, minwidth=80)
        elif self.warehouse_mode == "KMAT":
            self.tree.column("Time", width=80, minwidth=70)
            self.tree.column("Sales Order", width=110, minwidth=80)
            self.tree.column("Order", width=90, minwidth=70)
            self.tree.column("Material No.", width=100, minwidth=80)
            self.tree.column("Description", width=200, minwidth=120)
            self.tree.column("Quantity", width=80, minwidth=60)
            self.tree.column("UOM", width=60, minwidth=50)
            self.tree.column("Status", width=80, minwidth=60)
        elif self.warehouse_mode == "WIP":
            self.tree.column("Time", width=80, minwidth=70)
            self.tree.column("Sales Order", width=110, minwidth=80)
            self.tree.column("Order", width=90, minwidth=70)
            self.tree.column("Material No.", width=100, minwidth=80)
            self.tree.column("Description", width=200, minwidth=120)
            self.tree.column("Quantity", width=80, minwidth=60)
            self.tree.column("UOM", width=60, minwidth=50)
            self.tree.column("Status", width=80, minwidth=60)
        elif self.warehouse_mode == "Granules":
            self.tree.column("Time", width=90, minwidth=70)
            self.tree.column("Batch No.", width=130, minwidth=100)
            self.tree.column("Material No.", width=110, minwidth=80)
            self.tree.column("Description", width=240, minwidth=120)
            self.tree.column("Counted (KG)", width=100, minwidth=70)
            self.tree.column("Free Usable", width=100, minwidth=70)
            self.tree.column("Status", width=100, minwidth=80)
        else:
            self.tree.column("Time", width=90, minwidth=70)
            self.tree.column("Batch No.", width=130, minwidth=100)
            self.tree.column("Material", width=110, minwidth=80)
            self.tree.column("Shelf Location", width=100, minwidth=70)
            self.tree.column("Status", width=130, minwidth=100)

        self.tree.grid(row=1, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))

        v_scroll = ttk.Scrollbar(list_frame, orient=tk.VERTICAL, command=self.tree.yview)
        v_scroll.grid(row=1, column=1, sticky=(tk.N, tk.S))
        self.tree.configure(yscrollcommand=v_scroll.set)

        h_scroll = ttk.Scrollbar(list_frame, orient=tk.HORIZONTAL, command=self.tree.xview)
        h_scroll.grid(row=2, column=0, sticky=(tk.W, tk.E))
        self.tree.configure(xscrollcommand=h_scroll.set)

        self.tree.bind("<Button-3>", self.show_context_menu)

    # --- Button section ---

    def create_button_section(self):
        btn_frame = ttk.Frame(self.main_frame)
        btn_frame.grid(row=4, column=0, sticky=(tk.W, tk.E), pady=(8, 0))

        ttk.Button(btn_frame, text="Export / Backup", command=self.export_backup).grid(
            row=0, column=0, padx=(0, 10))
        ttk.Button(btn_frame, text="Settings", command=self.open_settings_dialog).grid(
            row=0, column=1, padx=(0, 10))
        ttk.Button(btn_frame, text="Fullscreen (F11)", command=self.toggle_fullscreen).grid(
            row=0, column=2, padx=(0, 10))
        ttk.Button(btn_frame, text="Quit", command=self.quit_app).grid(
            row=0, column=3, padx=(0, 10))

        self.switch_btn = ttk.Button(
            btn_frame, text="Switch Warehouse", command=self.switch_warehouse)
        self.switch_btn.grid(row=0, column=4)
        if not self.config.get("show_switch_warehouse", False):
            self.switch_btn.grid_remove()

    # --- Status bar ---

    def create_status_bar(self):
        self.status_var = tk.StringVar(value="Ready to scan...")
        status_bar = ttk.Label(
            self.root, textvariable=self.status_var,
            relief=tk.SUNKEN, anchor=tk.W, font=("Arial", 9))
        status_bar.grid(row=1, column=0, sticky=(tk.W, tk.E))

    # ------------------------------------------------------------------
    # Shortcuts
    # ------------------------------------------------------------------

    def bind_shortcuts(self):
        self.root.bind("<Control-z>", lambda e: self.undo_last_action())
        self.root.bind("<Control-Z>", lambda e: self.undo_last_action())
        self.root.bind("<Control-s>", lambda e: self._manual_save())
        self.root.bind("<Control-S>", lambda e: self._manual_save())
        self.root.bind("<Escape>", lambda e: self._reset_scan())
        self.root.bind("<F11>", lambda e: self.toggle_fullscreen())
        self.root.bind("<FocusIn>", self.ensure_scan_focus)

    # ------------------------------------------------------------------
    # Scan processing
    # ------------------------------------------------------------------

    def process_scan(self):
        """Process a scanned QR / barcode."""
        raw = self.scan_var.get().strip()
        if not raw:
            self.status_var.set("Please scan or enter a barcode / QR code.")
            return

        qr = self.parse_qr_code(raw)
        self.current_qr_data = qr
        charge = qr["charge"]

        self.logger.info(f"Scan received: raw='{raw}' -> charge='{charge}'")

        # Duplicate check
        if self._is_already_scanned(charge):
            messagebox.showwarning(
                "Already Scanned",
                f"Batch No. '{charge}' has already been scanned!\n\n"
                "Please check the list of scanned items.",
            )
            self._reset_scan()
            return

        if self.warehouse_mode == "Zert":
            row_data = self.suche_charge_zert(charge)
            if row_data is not None:
                self.show_found_zert(row_data, charge)
            else:
                self.show_not_found_dialog_zert(charge)
        elif self.warehouse_mode == "KMAT":
            order = raw.strip()
            if self._order_already_scanned(order, self.inventur_data_kmat, self.not_found_data_kmat):
                messagebox.showwarning(
                    "Already Scanned",
                    f"Order '{order}' has already been scanned!",
                )
                self._reset_scan()
                return

            row_data = self.suche_kmat(order)
            if row_data is not None:
                self.show_found_kmat(row_data, order)
            else:
                self.show_not_found_dialog_kmat(order)
        elif self.warehouse_mode == "WIP":
            order = raw.strip()
            if self._order_already_scanned(order, self.inventur_data_wip, self.not_found_data_wip):
                messagebox.showwarning(
                    "Already Scanned",
                    f"Order '{order}' has already been scanned!",
                )
                self._reset_scan()
                return

            row_data = self.suche_wip(order)
            if row_data is not None:
                self.show_found_wip(row_data, order)
            else:
                self.show_not_found_dialog_wip(order)
        elif self.warehouse_mode == "Granules":
            row_data = self.suche_granules(charge)
            if row_data is not None:
                self.show_found_granules(row_data, charge)
            else:
                self.show_not_found_dialog_granules(charge)
        else:
            row_data = self.suche_charge(charge)
            if row_data is not None:
                self.show_found_rolle(row_data, qr)
            else:
                self.show_not_found_dialog(charge, qr)

    def _is_already_scanned(self, charge):
        if self.warehouse_mode == "Zert":
            for item in self.inventur_data_zert + self.not_found_data_zert:
                if str(item.get("charge", "")) == str(charge):
                    return True
            return False
        elif self.warehouse_mode == "Granules":
            key = _batch_key(charge)
            for item in self.inventur_data_granules + self.not_found_data_granules:
                if _batch_key(item.get("charge", "")) == key:
                    return True
            return False
        else:
            for item in self.inventur_data + self.not_found_data:
                if str(item.get("charge", "")) == str(charge):
                    return True
            return False

    @staticmethod
    def _order_already_scanned(order, found_list, not_found_list):
        for item in found_list + not_found_list:
            if str(item.get("order", "")) == str(order):
                return True
        return False

    # ------------------------------------------------------------------
    # Show found roll (HALB)
    # ------------------------------------------------------------------

    def show_found_rolle(self, row_data, qr):
        """Display found roll data and show the input panel (HALB mode)."""
        def _fmt_int(val):
            """Format numeric value as integer string (remove .0)."""
            if val == "" or val is None:
                return ""
            try:
                return str(int(float(val)))
            except (ValueError, TypeError):
                return str(val)

        def _fmt_float2(val):
            """Format float to 2 decimal places, empty if missing."""
            if val == "" or val is None:
                return ""
            try:
                return f"{float(val):.2f}"
            except (ValueError, TypeError):
                return str(val)

        # Compute Area (m2) from length m × width mm / 1000
        laenge_m = row_data.get("length m", None)
        breite_mm = row_data.get("width mm", None)
        if laenge_m not in (None, "") and breite_mm not in (None, ""):
            try:
                flache_val = f"{float(laenge_m) * float(breite_mm) / 1000:.2f}"
            except (ValueError, TypeError):
                flache_val = ""
        else:
            flache_val = ""

        self.current_scan = {
            "charge": qr["charge"],
            "lort_master": _fmt_int(row_data.get("Storage Location", "") or ""),
            "lort_qr": qr["lort"],
            "material": _fmt_int(row_data.get("Material", "") or ""),
            "kurztext": str(row_data.get("Material Description", "") or ""),
            "werk": _fmt_int(row_data.get("Plant", "") or ""),
            "lnge0": qr["lnge0"],
            "brte0": qr["brte0"],
            "lnge1": qr["lnge1"],
            "brte1": qr["brte1"],
            "lnge2": qr["lnge2"],
            "brte2": qr["brte2"],
            "flache": flache_val,
            "frei_verw": _fmt_float2(row_data.get("Unrestricted", "") or ""),
            "status": "found",
        }

        # Update display labels
        self.lbl_charge.config(text=self.current_scan["charge"])
        self.lbl_material.config(text=self.current_scan["material"])
        self.lbl_kurztext.config(text=self.current_scan["kurztext"])
        self.lbl_werk.config(text=self.current_scan["werk"])
        self.lbl_lort_master.config(text=self.current_scan["lort_master"])
        self.lbl_lort_qr.config(text=self.current_scan["lort_qr"])
        self.lbl_flache.config(text=str(self.current_scan["flache"]))
        self.lbl_frei.config(text=str(self.current_scan["frei_verw"]))
        self.lbl_lnge0.config(text=str(self.current_scan["lnge0"]))
        self.lbl_brte0.config(text=str(self.current_scan["brte0"]))
        self.lbl_lnge1.config(text=str(self.current_scan["lnge1"]))
        self.lbl_brte1.config(text=str(self.current_scan["brte1"]))
        self.lbl_lnge2.config(text=str(self.current_scan["lnge2"]))
        self.lbl_brte2.config(text=str(self.current_scan["brte2"]))

        self.current_frame.config(text="ROLL FOUND")
        self.current_frame.grid()
        self._show_halb_widgets()

        # Reset input fields
        self.fach_var.set("")
        self.brte_meas_var.set("")
        self.remarks_var.set("")
        self.input_widgets["fach_entry"].focus_set()

        self.scan_var.set("")
        self.status_var.set(f"Roll found: {self.current_scan['kurztext']} | Enter Shelf Location")

    # ------------------------------------------------------------------
    # Show found Zert charge
    # ------------------------------------------------------------------

    def show_found_zert(self, row_data, charge):
        """Display found Zert charge data and show Zert input panel."""
        def _s(val):
            if val is None:
                return ""
            s = str(val)
            return "" if s.lower() == "nan" else s

        self.current_scan = {
            "charge": charge,
            "material": _s(row_data.get("Material Number", "")),
            "kurztext": _s(row_data.get("Material Description", "")),
            "mart": _s(row_data.get("MArt", "")),
            "werk": _s(row_data.get("Plant", "")),
            "lort": _s(row_data.get("Storage Location", "")),
            "bme": _s(row_data.get("Base Unit of Measure", "")),
            "frei_verw": _s(row_data.get("Unrestricted", "")),
            "laenge_mm": _s(row_data.get("Length", "")),
            "breite_mm": _s(row_data.get("Width", "")),
            "adv": _s(row_data.get("ADV", "")),
            "status": "found",
            "_mode": "zert",
        }

        # Update Zert info labels
        self.zert_lbl_charge.config(text=self.current_scan["charge"])
        self.zert_lbl_material.config(text=self.current_scan["material"])
        self.zert_lbl_kurztext.config(text=self.current_scan["kurztext"])
        self.zert_lbl_mart.config(text=self.current_scan["mart"])
        self.zert_lbl_werk.config(text=self.current_scan["werk"])
        self.zert_lbl_lort.config(text=self.current_scan["lort"])
        self.zert_lbl_bme.config(text=self.current_scan["bme"])
        self.zert_lbl_frei.config(text=self.current_scan["frei_verw"])
        self.zert_lbl_laenge.config(text=self.current_scan["laenge_mm"])
        self.zert_lbl_breite.config(text=self.current_scan["breite_mm"])
        self.zert_lbl_adv.config(text=self.current_scan["adv"])
        # Show BME hint next to menge entry
        self.zert_lbl_bme_input.config(text=self.current_scan["bme"])

        self.current_frame.config(text="BATCH FOUND (Zert)")
        self.current_frame.grid()
        self._show_zert_widgets()

        # Reset Zert input fields
        self.menge_var.set("")
        self.zert_remarks_var.set("")
        self.zert_input_widgets["menge_entry"].focus_set()

        self.scan_var.set("")
        self.status_var.set(
            f"Batch found: {self.current_scan['kurztext']} | Enter quantity")

    # ------------------------------------------------------------------
    # Show found KMAT item
    # ------------------------------------------------------------------

    def show_found_kmat(self, row_data, order):
        """Display found KMAT data and show KMAT input panel."""
        def _s(val):
            if val is None:
                return ""
            s = str(val)
            return "" if s.lower() == "nan" else s

        self.current_scan = {
            "sales_order": _s(row_data.get("Sales Order", "")),
            "order": order,
            "material": _s(row_data.get("Material Number", "")),
            "kurztext": _s(row_data.get("Material description", "")),
            "werk": _s(row_data.get("Plant", "")),
            "bme": _s(row_data.get("Unit of measure (=GMEIN)", "")),
            "status": "found",
            "_mode": "kmat",
        }

        self.kmat_lbl_sales_order.config(text=self.current_scan["sales_order"])
        self.kmat_lbl_order.config(text=self.current_scan["order"])
        self.kmat_lbl_material.config(text=self.current_scan["material"])
        self.kmat_lbl_kurztext.config(text=self.current_scan["kurztext"])
        self.kmat_lbl_werk.config(text=self.current_scan["werk"])
        self.kmat_lbl_bme.config(text=self.current_scan["bme"])
        self.kmat_lbl_bme_input.config(text=self.current_scan["bme"])

        self.current_frame.config(text="KMAT FOUND")
        self.current_frame.grid()
        self._show_kmat_widgets()

        self.kmat_menge_var.set("")
        self.kmat_remarks_var.set("")
        self.kmat_input_widgets["menge_entry"].focus_set()

        self.scan_var.set("")
        self.status_var.set(f"KMAT found: {self.current_scan['kurztext']} | Enter quantity")

    # ------------------------------------------------------------------
    # Show found WIP item
    # ------------------------------------------------------------------

    def show_found_wip(self, row_data, order):
        """Display found WIP data and show WIP input panel."""
        def _s(val):
            if val is None:
                return ""
            s = str(val)
            return "" if s.lower() == "nan" else s

        self.current_scan = {
            "sales_order": _s(row_data.get("Sales Order", "")),
            "order": order,
            "material": _s(row_data.get("Material Number", "")),
            "kurztext": _s(row_data.get("Material description", "")),
            "werk": _s(row_data.get("Plant", "")),
            "bme": _s(row_data.get("Unit of measure (=GMEIN)", "")),
            "order_qty": _s(row_data.get("Order quantity (GMEIN)", "")),
            "basic_finish": _s(row_data.get("Basic finish date", "")),
            "status": "found",
            "_mode": "wip",
        }

        self.wip_lbl_sales_order.config(text=self.current_scan["sales_order"])
        self.wip_lbl_order.config(text=self.current_scan["order"])
        self.wip_lbl_material.config(text=self.current_scan["material"])
        self.wip_lbl_kurztext.config(text=self.current_scan["kurztext"])
        self.wip_lbl_werk.config(text=self.current_scan["werk"])
        self.wip_lbl_bme.config(text=self.current_scan["bme"])
        self.wip_lbl_order_qty.config(text=self.current_scan["order_qty"])
        self.wip_lbl_basic_finish.config(text=self.current_scan["basic_finish"])

        self.current_frame.config(text="WIP FOUND")
        self.current_frame.grid()
        self._show_wip_widgets()

        self.wip_menge_var.set("")
        self.wip_remarks_var.set("")
        self.wip_input_widgets["menge_entry"].focus_set()

        self.scan_var.set("")
        self.status_var.set(f"WIP found: {self.current_scan['kurztext']} | Enter quantity")

    # ------------------------------------------------------------------
    # Not found dialog (HALB)
    # ------------------------------------------------------------------

    def show_not_found_dialog(self, charge, qr):
        """Show dialog for rolls not in master table (HALB)."""
        dlg = NotFoundDialogHALB(self.root, charge, qr)

        if dlg.result:
            data = dlg.result.copy()
            data["zeitstempel"] = datetime.now().strftime("%d.%m.%Y %H:%M:%S")

            self.not_found_data.append(data)

            self.undo_stack.append(("add_not_found", data.copy()))
            if len(self.undo_stack) > 50:
                self.undo_stack.pop(0)

            if self.config.get("auto_save", True):
                self.save_halb_excel()

            self.update_list()
            total = len(self.inventur_data) + len(self.not_found_data)
            self.status_var.set(
                f"Not-found roll saved. Total: {total} "
                f"({len(self.inventur_data)} found, {len(self.not_found_data)} not found)"
            )
            self.logger.info(f"Not-found roll saved: {data['charge']}")
        else:
            self.logger.info("Not-found dialog cancelled")

        self._reset_scan()

    # ------------------------------------------------------------------
    # Not found dialog (Zert)
    # ------------------------------------------------------------------

    def show_not_found_dialog_zert(self, charge):
        """Show dialog for Zert charges not in master table."""
        dlg = NotFoundDialogZert(self.root, charge)

        if dlg.result:
            data = dlg.result.copy()
            data["zeitstempel"] = datetime.now().strftime("%d.%m.%Y %H:%M:%S")

            self.not_found_data_zert.append(data)

            self.undo_stack.append(("add_not_found_zert", data.copy()))
            if len(self.undo_stack) > 50:
                self.undo_stack.pop(0)

            if self.config.get("auto_save", True):
                self.save_zert_excel()

            self.update_list()
            total = len(self.inventur_data_zert) + len(self.not_found_data_zert)
            self.status_var.set(
                f"Not-found batch saved. Total: {total} "
                f"({len(self.inventur_data_zert)} found, {len(self.not_found_data_zert)} not found)"
            )
            self.logger.info(f"Not-found Zert batch saved: {data['charge']}")
        else:
            self.logger.info("Zert not-found dialog cancelled")

        self._reset_scan()

    # ------------------------------------------------------------------
    # Not found dialog (KMAT)
    # ------------------------------------------------------------------

    def show_not_found_dialog_kmat(self, order):
        """Show dialog for KMAT Order not in master table."""
        dlg = NotFoundDialogKMAT(self.root, order)

        if dlg.result:
            data = dlg.result.copy()
            data["zeitstempel"] = datetime.now().strftime("%d.%m.%Y %H:%M:%S")

            self.not_found_data_kmat.append(data)

            self.undo_stack.append(("add_not_found_kmat", data.copy()))
            if len(self.undo_stack) > 50:
                self.undo_stack.pop(0)

            if self.config.get("auto_save", True):
                self.save_kmat_excel()

            self.update_list()
            total = len(self.inventur_data_kmat) + len(self.not_found_data_kmat)
            self.status_var.set(
                f"Not-found KMAT saved. Total: {total} "
                f"({len(self.inventur_data_kmat)} found, {len(self.not_found_data_kmat)} not found)"
            )
            self.logger.info(f"Not-found KMAT saved: {data['order']}")
        else:
            self.logger.info("KMAT not-found dialog cancelled")

        self._reset_scan()

    # ------------------------------------------------------------------
    # Not found dialog (WIP)
    # ------------------------------------------------------------------

    def show_not_found_dialog_wip(self, order):
        """Show dialog for WIP Order not in master table."""
        dlg = NotFoundDialogWIP(self.root, order)

        if dlg.result:
            data = dlg.result.copy()
            data["zeitstempel"] = datetime.now().strftime("%d.%m.%Y %H:%M:%S")

            self.not_found_data_wip.append(data)

            self.undo_stack.append(("add_not_found_wip", data.copy()))
            if len(self.undo_stack) > 50:
                self.undo_stack.pop(0)

            if self.config.get("auto_save", True):
                self.save_wip_excel()

            self.update_list()
            total = len(self.inventur_data_wip) + len(self.not_found_data_wip)
            self.status_var.set(
                f"Not-found WIP saved. Total: {total} "
                f"({len(self.inventur_data_wip)} found, {len(self.not_found_data_wip)} not found)"
            )
            self.logger.info(f"Not-found WIP saved: {data['order']}")
        else:
            self.logger.info("WIP not-found dialog cancelled")

        self._reset_scan()

    # ------------------------------------------------------------------
    # Save current scan (HALB)
    # ------------------------------------------------------------------

    def save_current_scan(self, event=None):
        """Validate and save the current found roll (HALB)."""
        if not self.current_scan:
            return

        fach = self.fach_var.get().strip()
        if not fach:
            messagebox.showwarning(
                "Required Field",
                "Shelf Location is mandatory. Please enter a value.",
            )
            self.input_widgets["fach_entry"].focus_set()
            return

        brte_meas = self.brte_meas_var.get().strip()
        if not brte_meas:
            messagebox.showwarning(
                "Required Field",
                "Measured Width (mm) is mandatory. Please enter a value.",
            )
            self.input_widgets["brte_meas_entry"].focus_set()
            return

        self.current_scan["fach"] = fach
        self.current_scan["brte_meas"] = brte_meas
        self.current_scan["remarks"] = self.remarks_var.get().strip()
        self.current_scan["zeitstempel"] = datetime.now().strftime("%d.%m.%Y %H:%M:%S")

        self.inventur_data.append(self.current_scan.copy())

        self.undo_stack.append(("add_found", self.current_scan.copy()))
        if len(self.undo_stack) > 50:
            self.undo_stack.pop(0)

        if self.config.get("auto_save", True):
            self.save_halb_excel()

        self.update_list()

        total = len(self.inventur_data) + len(self.not_found_data)
        self.status_var.set(
            f"Roll saved. Total: {total} "
            f"({len(self.inventur_data)} found, {len(self.not_found_data)} not found)"
        )
        self.logger.info(f"Roll saved: {self.current_scan['charge']}")

        self._reset_scan()

    # ------------------------------------------------------------------
    # Save current scan (Zert)
    # ------------------------------------------------------------------

    def save_current_scan_zert(self, event=None):
        """Validate and save the current found Zert charge."""
        if not self.current_scan:
            return

        menge = self.menge_var.get().strip()
        if not menge:
            messagebox.showwarning(
                "Required Field",
                "Recorded Quantity is a required field.",
            )
            self.zert_input_widgets["menge_entry"].focus_set()
            return

        # Validate numeric
        try:
            float(menge.replace(",", "."))
        except ValueError:
            messagebox.showwarning(
                "Invalid Input",
                "Recorded Quantity must be a number.",
            )
            self.zert_input_widgets["menge_entry"].focus_set()
            return

        self.current_scan["menge"] = menge
        self.current_scan["remarks"] = self.zert_remarks_var.get().strip()
        self.current_scan["zeitstempel"] = datetime.now().strftime("%d.%m.%Y %H:%M:%S")

        self.inventur_data_zert.append(self.current_scan.copy())

        self.undo_stack.append(("add_found_zert", self.current_scan.copy()))
        if len(self.undo_stack) > 50:
            self.undo_stack.pop(0)

        if self.config.get("auto_save", True):
            self.save_zert_excel()

        self.update_list()

        total = len(self.inventur_data_zert) + len(self.not_found_data_zert)
        self.status_var.set(
            f"Batch saved. Total: {total} "
            f"({len(self.inventur_data_zert)} found, {len(self.not_found_data_zert)} not found)"
        )
        self.logger.info(f"Zert batch saved: {self.current_scan['charge']}")

        self._reset_scan()

    # ------------------------------------------------------------------
    # Save current scan (KMAT)
    # ------------------------------------------------------------------

    def save_current_scan_kmat(self, event=None):
        """Validate and save the current found KMAT item."""
        if not self.current_scan or self.current_scan.get("_mode") != "kmat":
            return

        menge = self.kmat_menge_var.get().strip()
        if not menge:
            messagebox.showwarning("Required Field", "Recorded Quantity is mandatory.")
            self.kmat_input_widgets["menge_entry"].focus_set()
            return

        self.current_scan["menge"] = menge
        self.current_scan["remarks"] = self.kmat_remarks_var.get().strip()
        self.current_scan["zeitstempel"] = datetime.now().strftime("%d.%m.%Y %H:%M:%S")

        self.inventur_data_kmat.append(self.current_scan.copy())

        self.undo_stack.append(("add_kmat", self.current_scan.copy()))
        if len(self.undo_stack) > 50:
            self.undo_stack.pop(0)

        if self.config.get("auto_save", True):
            self.save_kmat_excel()

        self.update_list()
        total = len(self.inventur_data_kmat) + len(self.not_found_data_kmat)
        self.status_var.set(
            f"KMAT saved. Total: {total} "
            f"({len(self.inventur_data_kmat)} found, {len(self.not_found_data_kmat)} not found)"
        )
        self.logger.info(f"KMAT scan saved: {self.current_scan['sales_order']}/{self.current_scan['order']}")

        self.current_frame.grid_remove()
        self._hide_all_scan_widgets()
        self.current_scan = None
        self._reset_scan()

    # ------------------------------------------------------------------
    # Save current scan (WIP)
    # ------------------------------------------------------------------

    def save_current_scan_wip(self, event=None):
        """Validate and save the current found WIP item."""
        if not self.current_scan or self.current_scan.get("_mode") != "wip":
            return

        menge = self.wip_menge_var.get().strip()
        if not menge:
            messagebox.showwarning("Required Field", "Recorded Quantity is mandatory.")
            self.wip_input_widgets["menge_entry"].focus_set()
            return

        self.current_scan["menge"] = menge
        self.current_scan["remarks"] = self.wip_remarks_var.get().strip()
        self.current_scan["zeitstempel"] = datetime.now().strftime("%d.%m.%Y %H:%M:%S")

        self.inventur_data_wip.append(self.current_scan.copy())

        self.undo_stack.append(("add_wip", self.current_scan.copy()))
        if len(self.undo_stack) > 50:
            self.undo_stack.pop(0)

        if self.config.get("auto_save", True):
            self.save_wip_excel()

        self.update_list()
        self.status_var.set(
            f"WIP saved. Total: {len(self.inventur_data_wip)} entries"
        )
        self.logger.info(f"WIP scan saved: {self.current_scan['sales_order']}/{self.current_scan['order']}")

        self.current_frame.grid_remove()
        self._hide_all_scan_widgets()
        self.current_scan = None
        self._reset_scan()

    # ------------------------------------------------------------------
    # Granules: show found / not found / save
    # ------------------------------------------------------------------

    def show_found_granules(self, row_data, charge):
        """Display found Granules batch and show the Granules panel."""
        def _s(val):
            if val is None:
                return ""
            s = str(val)
            return "" if s.lower() == "nan" else s

        self.current_scan = {
            # Keep the batch as written in the master table (with leading zeros)
            "charge": _s(row_data.get("Batch", "")) or charge,
            "material": _s(row_data.get("Material Number", "")),
            "kurztext": _s(row_data.get("Material Description", "")),
            "werk": _s(row_data.get("Plant", "")),
            "lort": _s(row_data.get("Storage Location", "")),
            "bme": _s(row_data.get("Base Unit of Measure", "")) or "KG",
            "frei_verw": _fmt_qty(_s(row_data.get("Unrestricted", ""))),
            "status": "found",
            "_mode": "granules",
        }

        self.gran_lbl_charge.config(text=self.current_scan["charge"])
        self.gran_lbl_material.config(text=self.current_scan["material"])
        self.gran_lbl_kurztext.config(text=self.current_scan["kurztext"])
        self.gran_lbl_frei.config(
            text=f"{self.current_scan['frei_verw']} {self.current_scan['bme']}".strip())

        self.current_frame.config(text="BATCH FOUND (Granules)")
        self.current_frame.grid()
        self._show_granules_widgets()

        self.gran_menge_var.set("")
        self.gran_remarks_var.set("")
        self.gran_input_widgets["menge_entry"].focus_set()

        self.scan_var.set("")
        self.status_var.set(f"Batch found: {self.current_scan['kurztext']} | Enter counted quantity (KG)")

    def show_not_found_dialog_granules(self, charge):
        """Show dialog for Granules batches not in master table."""
        dlg = NotFoundDialogGranules(self.root, charge)

        if dlg.result:
            data = dlg.result.copy()
            data["zeitstempel"] = datetime.now().strftime("%d.%m.%Y %H:%M:%S")

            self.not_found_data_granules.append(data)

            self.undo_stack.append(("add_not_found_granules", data.copy()))
            if len(self.undo_stack) > 50:
                self.undo_stack.pop(0)

            if self.config.get("auto_save", True):
                self.save_granules_excel()

            self.update_list()
            total = len(self.inventur_data_granules) + len(self.not_found_data_granules)
            self.status_var.set(
                f"Not-found batch saved. Total: {total} "
                f"({len(self.inventur_data_granules)} found, {len(self.not_found_data_granules)} not found)"
            )
            self.logger.info(f"Not-found Granules batch saved: {data['charge']}")
        else:
            self.logger.info("Granules not-found dialog cancelled")

        self._reset_scan()

    def save_current_scan_granules(self, event=None):
        """Validate and save the current found Granules batch."""
        if not self.current_scan or self.current_scan.get("_mode") != "granules":
            return

        menge = self.gran_menge_var.get().strip()
        if not menge:
            messagebox.showwarning("Required Field", "Counted Qty (KG) is mandatory.")
            self.gran_input_widgets["menge_entry"].focus_set()
            return
        if _to_float(menge) is None:
            messagebox.showwarning("Invalid Input", "Counted Qty (KG) must be a number.")
            self.gran_input_widgets["menge_entry"].focus_set()
            return

        self.current_scan["menge"] = menge
        self.current_scan["remarks"] = self.gran_remarks_var.get().strip()
        self.current_scan["zeitstempel"] = datetime.now().strftime("%d.%m.%Y %H:%M:%S")

        self.inventur_data_granules.append(self.current_scan.copy())

        self.undo_stack.append(("add_found_granules", self.current_scan.copy()))
        if len(self.undo_stack) > 50:
            self.undo_stack.pop(0)

        if self.config.get("auto_save", True):
            self.save_granules_excel()

        self.update_list()
        total = len(self.inventur_data_granules) + len(self.not_found_data_granules)
        self.status_var.set(
            f"Batch saved. Total: {total} "
            f"({len(self.inventur_data_granules)} found, {len(self.not_found_data_granules)} not found)"
        )
        self.logger.info(f"Granules batch saved: {self.current_scan['charge']}")

        self._reset_scan()

    # ------------------------------------------------------------------
    # Reset scan
    # ------------------------------------------------------------------

    def _reset_scan(self):
        self.current_scan = None
        self.current_qr_data = None
        self.scan_var.set("")
        self.fach_var.set("")
        self.brte_meas_var.set("")
        self.remarks_var.set("")
        if hasattr(self, "menge_var"):
            self.menge_var.set("")
        if hasattr(self, "zert_remarks_var"):
            self.zert_remarks_var.set("")
        if hasattr(self, "kmat_menge_var"):
            self.kmat_menge_var.set("")
        if hasattr(self, "kmat_remarks_var"):
            self.kmat_remarks_var.set("")
        if hasattr(self, "wip_menge_var"):
            self.wip_menge_var.set("")
        if hasattr(self, "wip_remarks_var"):
            self.wip_remarks_var.set("")
        if hasattr(self, "gran_menge_var"):
            self.gran_menge_var.set("")
        if hasattr(self, "gran_remarks_var"):
            self.gran_remarks_var.set("")
        self.current_frame.grid_remove()
        self.scan_entry.focus_set()
        self.status_var.set("Ready to scan...")

    # ------------------------------------------------------------------
    # List update
    # ------------------------------------------------------------------

    @staticmethod
    def _parse_ts(ts_str):
        """Parse a timestamp string to datetime for correct sort order.

        Handles the in-session format DD.MM.YYYY HH:MM:SS and the ISO format
        that pandas produces when reading Excel date cells (YYYY-MM-DD HH:MM:SS).
        Falls back to datetime.min so unparseable values sort last.
        """
        for fmt in ("%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.strptime(ts_str, fmt)
            except ValueError:
                continue
        return datetime.min

    def update_list(self):
        for item in self.tree.get_children():
            self.tree.delete(item)

        if self.warehouse_mode == "Zert":
            all_items = []
            for d in self.inventur_data_zert:
                all_items.append((d, "Found"))
            for d in self.not_found_data_zert:
                all_items.append((d, "Not Found"))

            all_items.sort(key=lambda x: self._parse_ts(x[0].get("zeitstempel", "")), reverse=True)

            for d, status in all_items:
                ts = d.get("zeitstempel", "")
                time_part = ts.split(" ")[1] if " " in ts else ts
                self.tree.insert(
                    "", "end",
                    values=(
                        time_part,
                        d.get("charge", ""),
                        d.get("material", ""),
                        d.get("kurztext", ""),
                        d.get("menge", ""),
                        d.get("bme", ""),
                        status,
                    ),
                )

            total = len(self.inventur_data_zert) + len(self.not_found_data_zert)
            self.count_label.config(
                text=f"{total} entries ({len(self.inventur_data_zert)} found, "
                     f"{len(self.not_found_data_zert)} not found)"
            )
        elif self.warehouse_mode == "KMAT":
            all_items = []
            for d in self.inventur_data_kmat:
                all_items.append((d, "Found"))
            for d in self.not_found_data_kmat:
                all_items.append((d, "Not Found"))

            all_items.sort(key=lambda x: self._parse_ts(x[0].get("zeitstempel", "")), reverse=True)

            for d, status in all_items:
                ts = d.get("zeitstempel", "")
                time_part = ts.split(" ")[1] if " " in ts else ts
                self.tree.insert(
                    "", "end",
                    values=(
                        time_part,
                        d.get("sales_order", ""),
                        d.get("order", ""),
                        d.get("material", ""),
                        d.get("kurztext", ""),
                        d.get("menge", ""),
                        d.get("bme", ""),
                        status,
                    ),
                )

            total = len(self.inventur_data_kmat) + len(self.not_found_data_kmat)
            self.count_label.config(
                text=f"{total} entries ({len(self.inventur_data_kmat)} found, "
                     f"{len(self.not_found_data_kmat)} not found)"
            )
        elif self.warehouse_mode == "WIP":
            all_items = []
            for d in self.inventur_data_wip:
                all_items.append((d, "Found"))
            for d in self.not_found_data_wip:
                all_items.append((d, "Not Found"))

            all_items.sort(key=lambda x: self._parse_ts(x[0].get("zeitstempel", "")), reverse=True)

            for d, status in all_items:
                ts = d.get("zeitstempel", "")
                time_part = ts.split(" ")[1] if " " in ts else ts
                self.tree.insert(
                    "", "end",
                    values=(
                        time_part,
                        d.get("sales_order", ""),
                        d.get("order", ""),
                        d.get("material", ""),
                        d.get("kurztext", ""),
                        d.get("menge", ""),
                        d.get("bme", ""),
                        status,
                    ),
                )

            total = len(self.inventur_data_wip) + len(self.not_found_data_wip)
            self.count_label.config(
                text=f"{total} entries ({len(self.inventur_data_wip)} found, "
                     f"{len(self.not_found_data_wip)} not found)"
            )
        elif self.warehouse_mode == "Granules":
            all_items = []
            for d in self.inventur_data_granules:
                all_items.append((d, "Found"))
            for d in self.not_found_data_granules:
                all_items.append((d, "NOT FOUND"))

            all_items.sort(key=lambda x: self._parse_ts(x[0].get("zeitstempel", "")), reverse=True)

            for d, status in all_items:
                ts = d.get("zeitstempel", "")
                time_part = ts.split(" ")[1] if " " in ts else ts
                self.tree.insert(
                    "", "end",
                    values=(
                        time_part,
                        d.get("charge", ""),
                        d.get("material", ""),
                        d.get("kurztext", ""),
                        d.get("menge", ""),
                        d.get("frei_verw", ""),
                        status,
                    ),
                )

            total = len(self.inventur_data_granules) + len(self.not_found_data_granules)
            self.count_label.config(
                text=f"{total} batches ({len(self.inventur_data_granules)} found, "
                     f"{len(self.not_found_data_granules)} not found)"
            )
        else:
            all_items = []
            for d in self.inventur_data:
                all_items.append((d, "Found"))
            for d in self.not_found_data:
                all_items.append((d, "Not Found"))

            all_items.sort(key=lambda x: self._parse_ts(x[0].get("zeitstempel", "")), reverse=True)

            for d, status in all_items:
                ts = d.get("zeitstempel", "")
                time_part = ts.split(" ")[1] if " " in ts else ts
                item_id = self.tree.insert(
                    "", "end",
                    values=(
                        time_part,
                        d.get("charge", ""),
                        d.get("material", ""),
                        d.get("fach", ""),
                        status,
                    ),
                )
                if status == "Not Found":
                    self.tree.set(item_id, "Status", "NOT FOUND")
                else:
                    self.tree.set(item_id, "Status", "Found")

            total = len(self.inventur_data) + len(self.not_found_data)
            self.count_label.config(
                text=f"{total} rolls ({len(self.inventur_data)} found, "
                     f"{len(self.not_found_data)} not found)"
            )

    # ------------------------------------------------------------------
    # Excel save (HALB)
    # ------------------------------------------------------------------

    INVENTORY_HEADERS = [
        "Timestamp",
        "Plant",              # was: Werk
        "Location (Master)",  # was: LÖrt (Master)
        "Location (QR)",      # was: LÖrt (QR)
        "Material No.",       # was: Material
        "Description",        # was: Materialkurztext
        "Batch No.",          # was: Charge
        "Length S0 (mm)",     # was: Lnge
        "Width S0 (mm)",      # was: Brte
        "Length S1 (mm)",     # was: Lnge1
        "Width S1 (mm)",      # was: Brte1
        "Length S2 (mm)",     # was: Lnge2
        "Width S2 (mm)",      # was: Brte2
        "Area (m2)",          # was: Flache
        "Free Usable",        # was: +Frei verw.
        "Shelf Location",     # was: Shelf (Fach)
        "Measured Width (mm)",  # control value entered during scan
        "Remarks",
    ]

    def _row_from_item(self, d):
        def _clean(v):
            if v is None:
                return ""
            s = str(v)
            if s.lower() == "nan":
                return ""
            return s

        return [
            _clean(d.get("zeitstempel", "")),
            _clean(d.get("werk", "")),
            _clean(d.get("lort_master", "")),
            _clean(d.get("lort_qr", "")),
            _clean(d.get("material", "")),
            _clean(d.get("kurztext", "")),
            _clean(d.get("charge", "")),
            d.get("lnge0", ""),
            d.get("brte0", ""),
            d.get("lnge1", ""),
            d.get("brte1", ""),
            d.get("lnge2", ""),
            d.get("brte2", ""),
            _clean(d.get("flache", "")),
            _clean(d.get("frei_verw", "")),
            _clean(d.get("fach", "")),
            _clean(d.get("brte_meas", "")),
            _clean(d.get("remarks", "")),
        ]

    def _format_charge_col(self, ws, data_count):
        """Format Batch No. column (col 7 = G) as text."""
        charge_col = get_column_letter(7)
        for r in range(2, data_count + 2):
            ws[f"{charge_col}{r}"].number_format = "@"

    def save_halb_excel(self):
        """Write Inventory_<mode>.xlsx (HALB-type modes) with Inventory and Not_Found sheets."""
        try:
            wb = Workbook()
            if "Sheet" in wb.sheetnames:
                wb.remove(wb["Sheet"])

            # Inventory sheet
            ws_inv = wb.create_sheet("Inventory")
            ws_inv.append(self.INVENTORY_HEADERS)
            for d in self.inventur_data:
                ws_inv.append(self._row_from_item(d))
            self._format_charge_col(ws_inv, len(self.inventur_data))

            # Not_Found sheet
            ws_nf = wb.create_sheet("Not_Found")
            ws_nf.append(self.INVENTORY_HEADERS)
            for d in self.not_found_data:
                ws_nf.append(self._row_from_item(d))
            self._format_charge_col(ws_nf, len(self.not_found_data))

            # Ensure export dir exists
            self.export_path.mkdir(parents=True, exist_ok=True)
            wb.save(self.inventur_path)
            self.logger.info(f"Excel saved: {self.inventur_path}")

        except Exception as e:
            messagebox.showerror("Save Error", f"Error saving Excel file:\n{e}")
            self.logger.error(f"Error saving Excel: {e}")

    # ------------------------------------------------------------------
    # Excel save (Zert)
    # ------------------------------------------------------------------

    ZERT_HEADERS = [
        "Timestamp",
        "Material No.",
        "Description",
        "Mat. Type",
        "Plant",
        "Location",
        "Batch No.",
        "UOM",
        "Free Usable",
        "Length (mm)",
        "Width (mm)",
        "ADV",
        "Recorded Quantity",
        "Remarks",
    ]

    def _row_from_zert_item(self, d):
        def _clean(v):
            if v is None:
                return ""
            s = str(v)
            if s.lower() == "nan":
                return ""
            return s

        return [
            _clean(d.get("zeitstempel", "")),
            _clean(d.get("material", "")),
            _clean(d.get("kurztext", "")),
            _clean(d.get("mart", "")),
            _clean(d.get("werk", "")),
            _clean(d.get("lort", "")),
            _clean(d.get("charge", "")),
            _clean(d.get("bme", "")),
            _clean(d.get("frei_verw", "")),
            _clean(d.get("laenge_mm", "")),
            _clean(d.get("breite_mm", "")),
            _clean(d.get("adv", "")),
            _clean(d.get("menge", "")),
            _clean(d.get("remarks", "")),
        ]

    def save_zert_excel(self):
        """Write Inventory_Zert.xlsx with Inventory and Not_Found sheets."""
        try:
            wb = Workbook()
            if "Sheet" in wb.sheetnames:
                wb.remove(wb["Sheet"])

            ws_inv = wb.create_sheet("Inventory")
            ws_inv.append(self.ZERT_HEADERS)
            for d in self.inventur_data_zert:
                ws_inv.append(self._row_from_zert_item(d))

            ws_nf = wb.create_sheet("Not_Found")
            ws_nf.append(self.ZERT_HEADERS)
            for d in self.not_found_data_zert:
                ws_nf.append(self._row_from_zert_item(d))

            self.export_zert_path.mkdir(parents=True, exist_ok=True)
            wb.save(self.inventur_zert_path)
            self.logger.info(f"Zert Excel saved: {self.inventur_zert_path}")

        except Exception as e:
            messagebox.showerror("Save Error", f"Error saving Zert Excel file:\n{e}")
            self.logger.error(f"Error saving Zert Excel: {e}")

    # ------------------------------------------------------------------
    # Excel save (KMAT)
    # ------------------------------------------------------------------

    KMAT_HEADERS = [
        "Timestamp",
        "Plant",
        "Sales Order",
        "Order",
        "Material No.",
        "Description",
        "UOM",
        "Recorded Quantity",
        "Remarks",
    ]

    def _row_from_kmat_item(self, d):
        def _clean(v):
            if v is None:
                return ""
            s = str(v)
            return "" if s.lower() == "nan" else s

        return [
            _clean(d.get("zeitstempel", "")),
            _clean(d.get("werk", "")),
            _clean(d.get("sales_order", "")),
            _clean(d.get("order", "")),
            _clean(d.get("material", "")),
            _clean(d.get("kurztext", "")),
            _clean(d.get("bme", "")),
            _clean(d.get("menge", "")),
            _clean(d.get("remarks", "")),
        ]

    def save_kmat_excel(self):
        """Write Inventory_KMAT.xlsx with Inventory and Not_Found sheets."""
        try:
            wb = Workbook()
            if "Sheet" in wb.sheetnames:
                wb.remove(wb["Sheet"])

            ws_inv = wb.create_sheet("Inventory")
            ws_inv.append(self.KMAT_HEADERS)
            for d in self.inventur_data_kmat:
                ws_inv.append(self._row_from_kmat_item(d))

            ws_nf = wb.create_sheet("Not_Found")
            ws_nf.append(self.KMAT_HEADERS)
            for d in self.not_found_data_kmat:
                ws_nf.append(self._row_from_kmat_item(d))

            self.export_kmat_path.mkdir(parents=True, exist_ok=True)
            wb.save(self.inventur_kmat_path)
            self.logger.info(f"KMAT Excel saved: {self.inventur_kmat_path}")

        except Exception as e:
            messagebox.showerror("Save Error", f"Error saving KMAT Excel file:\n{e}")
            self.logger.error(f"Error saving KMAT Excel: {e}")

    # ------------------------------------------------------------------
    # Excel save (WIP)
    # ------------------------------------------------------------------

    WIP_HEADERS = [
        "Timestamp",
        "Plant",
        "Sales Order",
        "Order",
        "Material No.",
        "Description",
        "UOM",
        "Order Qty",
        "Recorded Quantity",
        "Remarks",
    ]

    def _row_from_wip_item(self, d):
        def _clean(v):
            if v is None:
                return ""
            s = str(v)
            return "" if s.lower() == "nan" else s

        return [
            _clean(d.get("zeitstempel", "")),
            _clean(d.get("werk", "")),
            _clean(d.get("sales_order", "")),
            _clean(d.get("order", "")),
            _clean(d.get("material", "")),
            _clean(d.get("kurztext", "")),
            _clean(d.get("bme", "")),
            _clean(d.get("order_qty", "")),
            _clean(d.get("menge", "")),
            _clean(d.get("remarks", "")),
        ]

    def save_wip_excel(self):
        """Write Inventory_WIP.xlsx with Inventory and Not_Found sheets."""
        try:
            wb = Workbook()
            if "Sheet" in wb.sheetnames:
                wb.remove(wb["Sheet"])

            ws_inv = wb.create_sheet("Inventory")
            ws_inv.append(self.WIP_HEADERS)
            for d in self.inventur_data_wip:
                ws_inv.append(self._row_from_wip_item(d))

            ws_nf = wb.create_sheet("Not_Found")
            ws_nf.append(self.WIP_HEADERS)
            for d in self.not_found_data_wip:
                ws_nf.append(self._row_from_wip_item(d))

            self.export_wip_path.mkdir(parents=True, exist_ok=True)
            wb.save(self.inventur_wip_path)
            self.logger.info(f"WIP Excel saved: {self.inventur_wip_path}")

        except Exception as e:
            messagebox.showerror("Save Error", f"Error saving WIP Excel file:\n{e}")
            self.logger.error(f"Error saving WIP Excel: {e}")

    # ------------------------------------------------------------------
    # Excel save (Granules)
    # ------------------------------------------------------------------

    GRANULES_HEADERS = [
        "Timestamp",
        "Plant",
        "Location",
        "Material No.",
        "Description",
        "Batch No.",
        "Free Usable",
        "UOM",
        "Counted Qty (KG)",
        "Difference (KG)",   # Counted Qty - Free Usable
        "Remarks",
    ]

    def _row_from_granules_item(self, d):
        def _clean(v):
            if v is None:
                return ""
            s = str(v)
            return "" if s.lower() == "nan" else s

        def _num(v):
            f = _to_float(v)
            if f is None:
                return _clean(v)
            return int(f) if f.is_integer() else f

        counted = _to_float(d.get("menge", ""))
        free = _to_float(d.get("frei_verw", ""))
        diff = _num(round(counted - free, 3)) if counted is not None and free is not None else ""

        return [
            _clean(d.get("zeitstempel", "")),
            _clean(d.get("werk", "")),
            _clean(d.get("lort", "")),
            _clean(d.get("material", "")),
            _clean(d.get("kurztext", "")),
            _clean(d.get("charge", "")),
            _num(d.get("frei_verw", "")),
            _clean(d.get("bme", "")),
            _num(d.get("menge", "")),
            diff,
            _clean(d.get("remarks", "")),
        ]

    def save_granules_excel(self):
        """Write Inventory_Granules.xlsx with Inventory and Not_Found sheets."""
        try:
            wb = Workbook()
            if "Sheet" in wb.sheetnames:
                wb.remove(wb["Sheet"])

            ws_inv = wb.create_sheet("Inventory")
            ws_inv.append(self.GRANULES_HEADERS)
            for d in self.inventur_data_granules:
                ws_inv.append(self._row_from_granules_item(d))

            ws_nf = wb.create_sheet("Not_Found")
            ws_nf.append(self.GRANULES_HEADERS)
            for d in self.not_found_data_granules:
                ws_nf.append(self._row_from_granules_item(d))

            # Batch No. (col F) as text so leading zeros are kept
            for ws, count in ((ws_inv, len(self.inventur_data_granules)),
                              (ws_nf, len(self.not_found_data_granules))):
                for r in range(2, count + 2):
                    ws[f"F{r}"].number_format = "@"

            self.export_granules_path.mkdir(parents=True, exist_ok=True)
            wb.save(self.inventur_granules_path)
            self.logger.info(f"Granules Excel saved: {self.inventur_granules_path}")

        except Exception as e:
            messagebox.showerror("Save Error", f"Error saving Granules Excel file:\n{e}")
            self.logger.error(f"Error saving Granules Excel: {e}")

    # ------------------------------------------------------------------
    # Load existing session (HALB)
    # ------------------------------------------------------------------

    def load_existing_halb(self):
        """Resume a previous session by loading Inventory_<mode>.xlsx (HALB-type modes)."""
        if not self.inventur_path.exists():
            return

        loaded = 0
        try:
            df_inv = pd.read_excel(
                self.inventur_path, sheet_name="Inventory", dtype={"Batch No.": str})
            for _, row in df_inv.iterrows():
                self.inventur_data.append(self._row_to_dict(row, status="found"))
                loaded += 1
        except Exception as e:
            self.logger.error(f"Error loading Inventory sheet: {e}")

        try:
            df_nf = pd.read_excel(
                self.inventur_path, sheet_name="Not_Found", dtype={"Batch No.": str})
            for _, row in df_nf.iterrows():
                self.not_found_data.append(self._row_to_dict(row, status="not_found"))
                loaded += 1
        except Exception:
            pass  # Sheet may not exist yet

        if loaded:
            self.update_list()
            self.status_var.set(
                f"Session resumed: {len(self.inventur_data)} found, "
                f"{len(self.not_found_data)} not found."
            )
            self.logger.info(f"Existing session loaded: {loaded} rows")

    def _row_to_dict(self, row, status):
        def _str(v):
            if pd.isna(v):
                return ""
            s = str(v)
            return "" if s.lower() == "nan" else s

        def _int_or_empty(v):
            try:
                return int(v)
            except (ValueError, TypeError):
                return 0

        return {
            "zeitstempel": _str(row.get("Timestamp", "")),
            "werk": _str(row.get("Plant", "")),
            "lort_master": _str(row.get("Location (Master)", "")),
            "lort_qr": _str(row.get("Location (QR)", "")),
            "material": _str(row.get("Material No.", "")),
            "kurztext": _str(row.get("Description", "")),
            "charge": _str(row.get("Batch No.", "")),
            "lnge0": _int_or_empty(row.get("Length S0 (mm)", 0)),
            "brte0": _int_or_empty(row.get("Width S0 (mm)", 0)),
            "lnge1": _int_or_empty(row.get("Length S1 (mm)", 0)),
            "brte1": _int_or_empty(row.get("Width S1 (mm)", 0)),
            "lnge2": _int_or_empty(row.get("Length S2 (mm)", 0)),
            "brte2": _int_or_empty(row.get("Width S2 (mm)", 0)),
            "flache": _str(row.get("Area (m2)", "")),
            "frei_verw": _str(row.get("Free Usable", "")),
            "fach": _str(row.get("Shelf Location", "")),
            "brte_meas": _str(row.get("Measured Width (mm)", "")),
            "remarks": _str(row.get("Remarks", "")),
            "status": status,
        }

    # ------------------------------------------------------------------
    # Load existing session (Zert)
    # ------------------------------------------------------------------

    def load_existing_zert(self):
        """Resume a previous Zert session by loading Inventory_Zert.xlsx."""
        if not self.inventur_zert_path.exists():
            return

        loaded = 0
        try:
            df_inv = pd.read_excel(
                self.inventur_zert_path, sheet_name="Inventory", dtype={"Batch No.": str, "Charge": str})
            for _, row in df_inv.iterrows():
                self.inventur_data_zert.append(self._row_to_zert_dict(row, status="found"))
                loaded += 1
        except Exception as e:
            self.logger.error(f"Error loading Zert Inventory sheet: {e}")

        try:
            df_nf = pd.read_excel(
                self.inventur_zert_path, sheet_name="Not_Found", dtype={"Batch No.": str, "Charge": str})
            for _, row in df_nf.iterrows():
                self.not_found_data_zert.append(self._row_to_zert_dict(row, status="not_found"))
                loaded += 1
        except Exception:
            pass

        if loaded:
            self.update_list()
            self.status_var.set(
                f"Zert session resumed: {len(self.inventur_data_zert)} found, "
                f"{len(self.not_found_data_zert)} not found."
            )
            self.logger.info(f"Existing Zert session loaded: {loaded} rows")

    def _row_to_zert_dict(self, row, status):
        def _str(v):
            try:
                if pd.isna(v):
                    return ""
            except (TypeError, ValueError):
                pass
            s = str(v)
            return "" if s.lower() == "nan" else s

        return {
            "zeitstempel": _str(row.get("Timestamp", "")),
            "material": _str(row.get("Material No.", "")),
            "kurztext": _str(row.get("Description", "")),
            "mart": _str(row.get("Mat. Type", "")),
            "werk": _str(row.get("Plant", "")),
            "lort": _str(row.get("Location", "")),
            "charge": _str(row.get("Batch No.", row.get("Charge", ""))),
            "bme": _str(row.get("UOM", "")),
            "frei_verw": _str(row.get("Free Usable", "")),
            "laenge_mm": _str(row.get("Length (mm)", "")),
            "breite_mm": _str(row.get("Width (mm)", "")),
            "adv": _str(row.get("ADV", "")),
            "menge": _str(row.get("Recorded Quantity", "")),
            "remarks": _str(row.get("Remarks", "")),
            "status": status,
        }

    # ------------------------------------------------------------------
    # Load existing session (KMAT)
    # ------------------------------------------------------------------

    def load_existing_kmat(self):
        """Resume a previous KMAT session by loading Inventory_KMAT.xlsx."""
        if not self.inventur_kmat_path.exists():
            return

        loaded = 0
        try:
            df_inv = pd.read_excel(
                self.inventur_kmat_path, sheet_name="Inventory",
                dtype={"Sales Order": str, "Order": str})
            for _, row in df_inv.iterrows():
                self.inventur_data_kmat.append(self._row_to_kmat_dict(row, status="found"))
                loaded += 1
        except Exception as e:
            self.logger.error(f"Error loading KMAT Inventory sheet: {e}")

        try:
            df_nf = pd.read_excel(
                self.inventur_kmat_path, sheet_name="Not_Found",
                dtype={"Sales Order": str, "Order": str})
            for _, row in df_nf.iterrows():
                self.not_found_data_kmat.append(self._row_to_kmat_dict(row, status="not_found"))
                loaded += 1
        except Exception:
            pass

        if loaded:
            self.update_list()
            self.status_var.set(
                f"KMAT session resumed: {len(self.inventur_data_kmat)} found, "
                f"{len(self.not_found_data_kmat)} not found."
            )
            self.logger.info(f"Existing KMAT session loaded: {loaded} rows")

    def _row_to_kmat_dict(self, row, status):
        def _str(v):
            try:
                if pd.isna(v):
                    return ""
            except (TypeError, ValueError):
                pass
            s = str(v)
            return "" if s.lower() == "nan" else s

        return {
            "zeitstempel": _str(row.get("Timestamp", "")),
            "werk": _str(row.get("Plant", "")),
            "sales_order": _str(row.get("Sales Order", "")),
            "order": _str(row.get("Order", "")),
            "material": _str(row.get("Material No.", "")),
            "kurztext": _str(row.get("Description", "")),
            "bme": _str(row.get("UOM", "")),
            "menge": _str(row.get("Recorded Quantity", "")),
            "remarks": _str(row.get("Remarks", "")),
            "status": status,
            "_mode": "kmat",
        }

    # ------------------------------------------------------------------
    # Load existing session (WIP)
    # ------------------------------------------------------------------

    def load_existing_wip(self):
        """Resume a previous WIP session by loading Inventory_WIP.xlsx."""
        if not self.inventur_wip_path.exists():
            return

        loaded = 0
        try:
            df_inv = pd.read_excel(
                self.inventur_wip_path, sheet_name="Inventory",
                dtype={"Sales Order": str, "Order": str})
            for _, row in df_inv.iterrows():
                self.inventur_data_wip.append(self._row_to_wip_dict(row, status="found"))
                loaded += 1
        except Exception as e:
            self.logger.error(f"Error loading WIP Inventory sheet: {e}")

        try:
            df_nf = pd.read_excel(
                self.inventur_wip_path, sheet_name="Not_Found",
                dtype={"Sales Order": str, "Order": str})
            for _, row in df_nf.iterrows():
                self.not_found_data_wip.append(self._row_to_wip_dict(row, status="not_found"))
                loaded += 1
        except Exception:
            pass

        if loaded:
            self.update_list()
            self.status_var.set(
                f"WIP session resumed: {len(self.inventur_data_wip)} found, "
                f"{len(self.not_found_data_wip)} not found."
            )
            self.logger.info(f"Existing WIP session loaded: {loaded} rows")

    def _row_to_wip_dict(self, row, status):
        def _str(v):
            try:
                if pd.isna(v):
                    return ""
            except (TypeError, ValueError):
                pass
            s = str(v)
            return "" if s.lower() == "nan" else s

        return {
            "zeitstempel": _str(row.get("Timestamp", "")),
            "werk": _str(row.get("Plant", "")),
            "sales_order": _str(row.get("Sales Order", "")),
            "order": _str(row.get("Order", "")),
            "material": _str(row.get("Material No.", "")),
            "kurztext": _str(row.get("Description", "")),
            "bme": _str(row.get("UOM", "")),
            "order_qty": _str(row.get("Order Qty", "")),
            "menge": _str(row.get("Recorded Quantity", "")),
            "remarks": _str(row.get("Remarks", "")),
            "status": status,
            "_mode": "wip",
        }

    # ------------------------------------------------------------------
    # Load existing session (Granules)
    # ------------------------------------------------------------------

    def load_existing_granules(self):
        """Resume a previous Granules session by loading Inventory_Granules.xlsx."""
        if not self.inventur_granules_path.exists():
            return

        loaded = 0
        try:
            df_inv = pd.read_excel(
                self.inventur_granules_path, sheet_name="Inventory",
                dtype={"Batch No.": str, "Material No.": str, "Plant": str, "Location": str, "Remarks": str})
            for _, row in df_inv.iterrows():
                self.inventur_data_granules.append(self._row_to_granules_dict(row, status="found"))
                loaded += 1
        except Exception as e:
            self.logger.error(f"Error loading Granules Inventory sheet: {e}")

        try:
            df_nf = pd.read_excel(
                self.inventur_granules_path, sheet_name="Not_Found",
                dtype={"Batch No.": str, "Material No.": str, "Plant": str, "Location": str, "Remarks": str})
            for _, row in df_nf.iterrows():
                self.not_found_data_granules.append(self._row_to_granules_dict(row, status="not_found"))
                loaded += 1
        except Exception:
            pass

        if loaded:
            self.update_list()
            self.status_var.set(
                f"Granules session resumed: {len(self.inventur_data_granules)} found, "
                f"{len(self.not_found_data_granules)} not found."
            )
            self.logger.info(f"Existing Granules session loaded: {loaded} rows")

    def _row_to_granules_dict(self, row, status):
        def _str(v):
            try:
                if pd.isna(v):
                    return ""
            except (TypeError, ValueError):
                pass
            s = str(v)
            return "" if s.lower() == "nan" else s

        return {
            "zeitstempel": _str(row.get("Timestamp", "")),
            "werk": _str(row.get("Plant", "")),
            "lort": _str(row.get("Location", "")),
            "material": _str(row.get("Material No.", "")),
            "kurztext": _str(row.get("Description", "")),
            "charge": _str(row.get("Batch No.", "")),
            "frei_verw": _fmt_qty(_str(row.get("Free Usable", ""))),
            "bme": _str(row.get("UOM", "")),
            "menge": _fmt_qty(_str(row.get("Counted Qty (KG)", ""))),
            "remarks": _str(row.get("Remarks", "")),
            "status": status,
            "_mode": "granules",
        }

    # ------------------------------------------------------------------
    # Export / Backup
    # ------------------------------------------------------------------

    def export_backup(self):
        """Save a timestamped backup of the current Excel file."""
        try:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

            if self.warehouse_mode == "Zert":
                self.save_zert_excel()
                backup_dir = self.export_zert_path / "backups"
                backup_dir.mkdir(parents=True, exist_ok=True)
                backup_file = backup_dir / f"Inventory_Zert_Backup_{timestamp}.xlsx"
                source_file = self.inventur_zert_path
                if source_file.exists():
                    shutil.copy2(source_file, backup_file)
                    messagebox.showinfo(
                        "Backup Created",
                        f"Backup saved to:\n{backup_file}",
                    )
                    self.logger.info(f"Zert backup created: {backup_file}")
                else:
                    messagebox.showwarning("Warning", "No Zert data file to backup yet.")
            elif self.warehouse_mode == "KMAT":
                self.save_kmat_excel()
                backup_dir = self.export_kmat_path / "backups"
                backup_dir.mkdir(parents=True, exist_ok=True)
                backup_file = backup_dir / f"Inventory_KMAT_Backup_{timestamp}.xlsx"
                if self.inventur_kmat_path.exists():
                    shutil.copy2(self.inventur_kmat_path, backup_file)
                    messagebox.showinfo(
                        "Backup Created",
                        f"Backup saved to:\n{backup_file}",
                    )
                    self.logger.info(f"KMAT backup created: {backup_file}")
                else:
                    messagebox.showwarning("Warning", "No KMAT data file to backup yet.")
            elif self.warehouse_mode == "WIP":
                self.save_wip_excel()
                backup_dir = self.export_wip_path / "backups"
                backup_dir.mkdir(parents=True, exist_ok=True)
                backup_file = backup_dir / f"Inventory_WIP_Backup_{timestamp}.xlsx"
                source_file = self.inventur_wip_path
                if source_file.exists():
                    shutil.copy2(source_file, backup_file)
                    messagebox.showinfo(
                        "Backup Created",
                        f"Backup saved to:\n{backup_file}",
                    )
                    self.logger.info(f"WIP backup created: {backup_file}")
                else:
                    messagebox.showwarning("Warning", "No WIP data file to backup yet.")
            elif self.warehouse_mode == "Granules":
                self.save_granules_excel()
                backup_dir = self.export_granules_path / "backups"
                backup_dir.mkdir(parents=True, exist_ok=True)
                backup_file = backup_dir / f"Inventory_Granules_Backup_{timestamp}.xlsx"
                if self.inventur_granules_path.exists():
                    shutil.copy2(self.inventur_granules_path, backup_file)
                    messagebox.showinfo(
                        "Backup Created",
                        f"Backup saved to:\n{backup_file}",
                    )
                    self.logger.info(f"Granules backup created: {backup_file}")
                else:
                    messagebox.showwarning("Warning", "No Granules data file to backup yet.")
            else:
                self.save_halb_excel()
                backup_dir = self.export_path / "backups"
                backup_dir.mkdir(parents=True, exist_ok=True)
                backup_file = backup_dir / f"Inventory_{self.halb_mode}_Backup_{timestamp}.xlsx"
                if self.inventur_path.exists():
                    shutil.copy2(self.inventur_path, backup_file)
                    messagebox.showinfo(
                        "Backup Created",
                        f"Backup saved to:\n{backup_file}",
                    )
                    self.logger.info(f"Backup created: {backup_file}")
                else:
                    messagebox.showwarning("Warning", "No data file to backup yet.")
        except Exception as e:
            messagebox.showerror("Export Error", f"Error creating backup:\n{e}")
            self.logger.error(f"Backup error: {e}")

    # ------------------------------------------------------------------
    # Undo
    # ------------------------------------------------------------------

    def undo_last_action(self):
        if not self.undo_stack:
            self.status_var.set("Nothing to undo")
            return

        action, data = self.undo_stack.pop()
        charge = data.get("charge", "")

        if action == "add_found":
            self.inventur_data = [
                d for d in self.inventur_data if d.get("charge") != charge]
            if self.config.get("auto_save", True):
                self.save_halb_excel()
        elif action == "add_not_found":
            self.not_found_data = [
                d for d in self.not_found_data if d.get("charge") != charge]
            if self.config.get("auto_save", True):
                self.save_halb_excel()
        elif action == "add_found_zert":
            self.inventur_data_zert = [
                d for d in self.inventur_data_zert if d.get("charge") != charge]
            if self.config.get("auto_save", True):
                self.save_zert_excel()
        elif action == "add_not_found_zert":
            self.not_found_data_zert = [
                d for d in self.not_found_data_zert if d.get("charge") != charge]
            if self.config.get("auto_save", True):
                self.save_zert_excel()
        elif action == "add_kmat":
            order = data.get("order", "")
            self.inventur_data_kmat = [
                d for d in self.inventur_data_kmat if d.get("order") != order]
            if self.config.get("auto_save", True):
                self.save_kmat_excel()
        elif action == "add_not_found_kmat":
            order = data.get("order", "")
            self.not_found_data_kmat = [
                d for d in self.not_found_data_kmat if d.get("order") != order]
            if self.config.get("auto_save", True):
                self.save_kmat_excel()
        elif action == "add_wip":
            order = data.get("order", "")
            self.inventur_data_wip = [
                d for d in self.inventur_data_wip if d.get("order") != order]
            if self.config.get("auto_save", True):
                self.save_wip_excel()
        elif action == "add_not_found_wip":
            order = data.get("order", "")
            self.not_found_data_wip = [
                d for d in self.not_found_data_wip if d.get("order") != order]
            if self.config.get("auto_save", True):
                self.save_wip_excel()
        elif action == "add_found_granules":
            self.inventur_data_granules = [
                d for d in self.inventur_data_granules if d.get("charge") != charge]
            if self.config.get("auto_save", True):
                self.save_granules_excel()
        elif action == "add_not_found_granules":
            self.not_found_data_granules = [
                d for d in self.not_found_data_granules if d.get("charge") != charge]
            if self.config.get("auto_save", True):
                self.save_granules_excel()

        self.update_list()
        self.status_var.set(f"Undo: removed entry for batch '{charge}'")
        self.logger.info(f"Undo {action}: {charge}")

    # ------------------------------------------------------------------
    # Context menu
    # ------------------------------------------------------------------

    def show_context_menu(self, event):
        selection = self.tree.selection()
        if not selection:
            return
        item_id = selection[0]

        menu = tk.Menu(self.root, tearoff=0)
        menu.add_command(label="Edit entry",
                         command=lambda: self._edit_entry(item_id))
        menu.add_separator()
        menu.add_command(label="Delete entry",
                         command=lambda: self._delete_entry(item_id))
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()

    def _edit_entry(self, item_id):
        """Open a dialog to edit correctable fields of a scanned entry."""
        values = self.tree.item(item_id, "values")
        if not values:
            return

        # Find the matching entry in the data list
        if self.warehouse_mode == "KMAT":
            order = values[2] if len(values) > 2 else ""
            entry = next(
                (d for d in self.inventur_data_kmat if d.get("order") == order),
                None)
            if entry is None:
                entry = next(
                    (d for d in self.not_found_data_kmat if d.get("order") == order),
                    None)
        else:
            charge = values[1] if len(values) > 1 else ""
            if self.warehouse_mode == "Zert":
                entry = next(
                    (d for d in self.inventur_data_zert if d.get("charge") == charge),
                    None)
                if entry is None:
                    entry = next(
                        (d for d in self.not_found_data_zert if d.get("charge") == charge),
                        None)
            elif self.warehouse_mode == "WIP":
                order = values[2] if len(values) > 2 else ""
                entry = next(
                    (d for d in self.inventur_data_wip if d.get("order") == order),
                    None)
                if entry is None:
                    entry = next(
                        (d for d in self.not_found_data_wip if d.get("order") == order),
                        None)
            elif self.warehouse_mode == "Granules":
                # Treeview may return numeric-looking batches as int -> compare by batch key
                key = _batch_key(charge)
                entry = next(
                    (d for d in self.inventur_data_granules + self.not_found_data_granules
                     if _batch_key(d.get("charge")) == key),
                    None)
            else:  # HALB-type modes
                entry = next(
                    (d for d in self.inventur_data if d.get("charge") == charge),
                    None)
                if entry is None:
                    entry = next(
                        (d for d in self.not_found_data if d.get("charge") == charge),
                        None)

        if entry is None:
            messagebox.showerror("Error", "Entry not found in data.")
            return

        # Build modal edit dialog
        dlg = tk.Toplevel(self.root)
        dlg.title("Edit Entry")
        dlg.resizable(False, False)
        dlg.grab_set()
        dlg.transient(self.root)

        frame = ttk.Frame(dlg, padding="20")
        frame.grid(row=0, column=0)

        row = 0

        if self.warehouse_mode in HALB_MODES:
            # Show identifier (read-only)
            ttk.Label(frame, text="Batch No.:", font=("Arial", 16, "bold")).grid(
                row=row, column=0, sticky=tk.W, pady=4, padx=(0, 10))
            ttk.Label(frame, text=entry.get("charge", ""), font=("Arial", 16)).grid(
                row=row, column=1, sticky=tk.W, pady=4)
            row += 1

            ttk.Label(frame, text="Shelf Location *:", font=("Arial", 16, "bold")).grid(
                row=row, column=0, sticky=tk.W, pady=4, padx=(0, 10))
            fach_var = tk.StringVar(value=entry.get("fach", ""))
            fach_entry = ttk.Entry(frame, textvariable=fach_var, font=("Arial", 18), width=20)
            fach_entry.grid(row=row, column=1, sticky=tk.W, pady=4)
            row += 1

            ttk.Label(frame, text="Measured Width (mm) *:", font=("Arial", 16, "bold")).grid(
                row=row, column=0, sticky=tk.W, pady=4, padx=(0, 10))
            brte_var = tk.StringVar(value=entry.get("brte_meas", ""))
            brte_entry = ttk.Entry(frame, textvariable=brte_var, font=("Arial", 18), width=10)
            brte_entry.grid(row=row, column=1, sticky=tk.W, pady=4)
            row += 1

            ttk.Label(frame, text="Remarks:", font=("Arial", 16, "bold")).grid(
                row=row, column=0, sticky=tk.W, pady=4, padx=(0, 10))
            remarks_var = tk.StringVar(value=entry.get("remarks", ""))
            remarks_entry = ttk.Entry(frame, textvariable=remarks_var, font=("Arial", 18), width=40)
            remarks_entry.grid(row=row, column=1, sticky=tk.W, pady=4)
            row += 1

            def save_edit():
                new_fach = fach_var.get().strip()
                new_brte = brte_var.get().strip()
                new_remarks = remarks_var.get().strip()
                if not new_fach:
                    messagebox.showerror("Validation", "Shelf Location is required.", parent=dlg)
                    return
                if not new_brte:
                    messagebox.showerror("Validation", "Measured Width is required.", parent=dlg)
                    return
                try:
                    int(new_brte)
                except ValueError:
                    messagebox.showerror("Validation", "Measured Width must be a number.", parent=dlg)
                    return
                entry["fach"] = new_fach
                entry["brte_meas"] = new_brte
                entry["remarks"] = new_remarks
                self.save_halb_excel()
                self.update_list()
                self.status_var.set(f"Entry updated: {entry.get('charge', '')}")
                dlg.destroy()

            fach_entry.focus_set()
            fach_entry.bind("<Return>", lambda e: brte_entry.focus_set())
            brte_entry.bind("<Return>", lambda e: remarks_entry.focus_set())
            remarks_entry.bind("<Return>", lambda e: save_edit())

        elif self.warehouse_mode == "WIP":
            ttk.Label(frame, text="Sales Order:", font=("Arial", 16, "bold")).grid(
                row=row, column=0, sticky=tk.W, pady=4, padx=(0, 10))
            ttk.Label(frame, text=entry.get("sales_order", ""), font=("Arial", 16)).grid(
                row=row, column=1, sticky=tk.W, pady=4)
            row += 1

            ttk.Label(frame, text="Order:", font=("Arial", 16, "bold")).grid(
                row=row, column=0, sticky=tk.W, pady=4, padx=(0, 10))
            ttk.Label(frame, text=entry.get("order", ""), font=("Arial", 16)).grid(
                row=row, column=1, sticky=tk.W, pady=4)
            row += 1

            ttk.Label(frame, text=f"Quantity * ({entry.get('bme', '')}):", font=("Arial", 16, "bold")).grid(
                row=row, column=0, sticky=tk.W, pady=4, padx=(0, 10))
            menge_var = tk.StringVar(value=entry.get("menge", ""))
            menge_entry = ttk.Entry(frame, textvariable=menge_var, font=("Arial", 18), width=15)
            menge_entry.grid(row=row, column=1, sticky=tk.W, pady=4)
            row += 1

            ttk.Label(frame, text="Remarks:", font=("Arial", 16, "bold")).grid(
                row=row, column=0, sticky=tk.W, pady=4, padx=(0, 10))
            remarks_var = tk.StringVar(value=entry.get("remarks", ""))
            remarks_entry = ttk.Entry(frame, textvariable=remarks_var, font=("Arial", 18), width=40)
            remarks_entry.grid(row=row, column=1, sticky=tk.W, pady=4)
            row += 1

            def save_edit():
                new_menge = menge_var.get().strip()
                new_remarks = remarks_var.get().strip()
                if not new_menge:
                    messagebox.showerror("Validation", "Quantity is required.", parent=dlg)
                    return
                try:
                    float(new_menge)
                except ValueError:
                    messagebox.showerror("Validation", "Quantity must be a number.", parent=dlg)
                    return
                entry["menge"] = new_menge
                entry["remarks"] = new_remarks
                self.save_wip_excel()
                self.update_list()
                self.status_var.set(f"Entry updated: {entry.get('sales_order', '')}/{entry.get('order', '')}")
                dlg.destroy()

            menge_entry.focus_set()
            menge_entry.bind("<Return>", lambda e: remarks_entry.focus_set())
            remarks_entry.bind("<Return>", lambda e: save_edit())

        elif self.warehouse_mode == "Granules":
            ttk.Label(frame, text="Batch No.:", font=("Arial", 16, "bold")).grid(
                row=row, column=0, sticky=tk.W, pady=4, padx=(0, 10))
            ttk.Label(frame, text=entry.get("charge", ""), font=("Arial", 16)).grid(
                row=row, column=1, sticky=tk.W, pady=4)
            row += 1

            ttk.Label(frame, text="Counted Qty (KG) *:", font=("Arial", 16, "bold")).grid(
                row=row, column=0, sticky=tk.W, pady=4, padx=(0, 10))
            menge_var = tk.StringVar(value=entry.get("menge", ""))
            menge_entry = ttk.Entry(frame, textvariable=menge_var, font=("Arial", 18), width=15)
            menge_entry.grid(row=row, column=1, sticky=tk.W, pady=4)
            row += 1

            ttk.Label(frame, text="Remarks:", font=("Arial", 16, "bold")).grid(
                row=row, column=0, sticky=tk.W, pady=4, padx=(0, 10))
            remarks_var = tk.StringVar(value=entry.get("remarks", ""))
            remarks_entry = ttk.Entry(frame, textvariable=remarks_var, font=("Arial", 18), width=40)
            remarks_entry.grid(row=row, column=1, sticky=tk.W, pady=4)
            row += 1

            def save_edit():
                new_menge = menge_var.get().strip()
                new_remarks = remarks_var.get().strip()
                if not new_menge:
                    messagebox.showerror("Validation", "Counted Qty (KG) is required.", parent=dlg)
                    return
                if _to_float(new_menge) is None:
                    messagebox.showerror("Validation", "Counted Qty (KG) must be a number.", parent=dlg)
                    return
                entry["menge"] = new_menge
                entry["remarks"] = new_remarks
                self.save_granules_excel()
                self.update_list()
                self.status_var.set(f"Entry updated: {entry.get('charge', '')}")
                dlg.destroy()

            menge_entry.focus_set()
            menge_entry.bind("<Return>", lambda e: remarks_entry.focus_set())
            remarks_entry.bind("<Return>", lambda e: save_edit())

        elif self.warehouse_mode == "Zert":
            ttk.Label(frame, text="Batch No.:", font=("Arial", 16, "bold")).grid(
                row=row, column=0, sticky=tk.W, pady=4, padx=(0, 10))
            ttk.Label(frame, text=entry.get("charge", ""), font=("Arial", 16)).grid(
                row=row, column=1, sticky=tk.W, pady=4)
            row += 1

            ttk.Label(frame, text=f"Quantity * ({entry.get('bme', '')}):", font=("Arial", 16, "bold")).grid(
                row=row, column=0, sticky=tk.W, pady=4, padx=(0, 10))
            menge_var = tk.StringVar(value=entry.get("menge", ""))
            menge_entry = ttk.Entry(frame, textvariable=menge_var, font=("Arial", 18), width=15)
            menge_entry.grid(row=row, column=1, sticky=tk.W, pady=4)
            row += 1

            ttk.Label(frame, text="Remarks:", font=("Arial", 16, "bold")).grid(
                row=row, column=0, sticky=tk.W, pady=4, padx=(0, 10))
            remarks_var = tk.StringVar(value=entry.get("remarks", ""))
            remarks_entry = ttk.Entry(frame, textvariable=remarks_var, font=("Arial", 18), width=40)
            remarks_entry.grid(row=row, column=1, sticky=tk.W, pady=4)
            row += 1

            def save_edit():
                new_menge = menge_var.get().strip()
                new_remarks = remarks_var.get().strip()
                if not new_menge:
                    messagebox.showerror("Validation", "Quantity is required.", parent=dlg)
                    return
                try:
                    float(new_menge)
                except ValueError:
                    messagebox.showerror("Validation", "Quantity must be a number.", parent=dlg)
                    return
                entry["menge"] = new_menge
                entry["remarks"] = new_remarks
                self.save_zert_excel()
                self.update_list()
                self.status_var.set(f"Entry updated: {entry.get('charge', '')}")
                dlg.destroy()

            menge_entry.focus_set()
            menge_entry.bind("<Return>", lambda e: remarks_entry.focus_set())
            remarks_entry.bind("<Return>", lambda e: save_edit())

        else:  # KMAT
            ttk.Label(frame, text="Sales Order:", font=("Arial", 16, "bold")).grid(
                row=row, column=0, sticky=tk.W, pady=4, padx=(0, 10))
            ttk.Label(frame, text=entry.get("sales_order", ""), font=("Arial", 16)).grid(
                row=row, column=1, sticky=tk.W, pady=4)
            row += 1

            ttk.Label(frame, text="Order:", font=("Arial", 16, "bold")).grid(
                row=row, column=0, sticky=tk.W, pady=4, padx=(0, 10))
            ttk.Label(frame, text=entry.get("order", ""), font=("Arial", 16)).grid(
                row=row, column=1, sticky=tk.W, pady=4)
            row += 1

            ttk.Label(frame, text=f"Quantity * ({entry.get('bme', '')}):", font=("Arial", 16, "bold")).grid(
                row=row, column=0, sticky=tk.W, pady=4, padx=(0, 10))
            menge_var = tk.StringVar(value=entry.get("menge", ""))
            menge_entry = ttk.Entry(frame, textvariable=menge_var, font=("Arial", 18), width=15)
            menge_entry.grid(row=row, column=1, sticky=tk.W, pady=4)
            row += 1

            ttk.Label(frame, text="Remarks:", font=("Arial", 16, "bold")).grid(
                row=row, column=0, sticky=tk.W, pady=4, padx=(0, 10))
            remarks_var = tk.StringVar(value=entry.get("remarks", ""))
            remarks_entry = ttk.Entry(frame, textvariable=remarks_var, font=("Arial", 18), width=40)
            remarks_entry.grid(row=row, column=1, sticky=tk.W, pady=4)
            row += 1

            def save_edit():
                new_menge = menge_var.get().strip()
                new_remarks = remarks_var.get().strip()
                if not new_menge:
                    messagebox.showerror("Validation", "Quantity is required.", parent=dlg)
                    return
                try:
                    float(new_menge)
                except ValueError:
                    messagebox.showerror("Validation", "Quantity must be a number.", parent=dlg)
                    return
                entry["menge"] = new_menge
                entry["remarks"] = new_remarks
                self.save_kmat_excel()
                self.update_list()
                self.status_var.set(
                    f"Entry updated: {entry.get('sales_order', '')}/{entry.get('order', '')}")
                dlg.destroy()

            menge_entry.focus_set()
            menge_entry.bind("<Return>", lambda e: remarks_entry.focus_set())
            remarks_entry.bind("<Return>", lambda e: save_edit())

        # Buttons row
        btn_frame = ttk.Frame(frame)
        btn_frame.grid(row=row, column=0, columnspan=2, pady=(12, 0))
        ttk.Button(btn_frame, text="Save", command=save_edit,
                   style="BigSave.TButton").grid(row=0, column=0, padx=(0, 10))
        ttk.Button(btn_frame, text="Cancel", command=dlg.destroy,
                   style="BigSave.TButton").grid(row=0, column=1)

        # Center dialog on parent
        dlg.update_idletasks()
        x = self.root.winfo_x() + (self.root.winfo_width() - dlg.winfo_width()) // 2
        y = self.root.winfo_y() + (self.root.winfo_height() - dlg.winfo_height()) // 2
        dlg.geometry(f"+{x}+{y}")
        dlg.wait_window()

    def _delete_entry(self, item_id):
        if not messagebox.askyesno("Confirm Delete",
                                   "Delete this entry permanently?"):
            return

        values = self.tree.item(item_id, "values")
        if len(values) >= 2:
            if self.warehouse_mode == "KMAT":
                # For KMAT: values = (Time, Sales Order, Order, ...)
                order = values[2] if len(values) > 2 else ""
                self.inventur_data_kmat = [
                    d for d in self.inventur_data_kmat if d.get("order") != order]
                self.not_found_data_kmat = [
                    d for d in self.not_found_data_kmat if d.get("order") != order]
                self.save_kmat_excel()
                self.update_list()
                self.status_var.set(f"Entry deleted: {order}")
                self.logger.info(f"KMAT entry deleted: {order}")
            elif self.warehouse_mode == "WIP":
                order = values[2] if len(values) > 2 else ""
                self.inventur_data_wip = [
                    d for d in self.inventur_data_wip if d.get("order") != order]
                self.not_found_data_wip = [
                    d for d in self.not_found_data_wip if d.get("order") != order]
                self.save_wip_excel()
                self.update_list()
                self.status_var.set(f"Entry deleted: {order}")
                self.logger.info(f"WIP entry deleted: {order}")
            else:
                charge = values[1]
                if self.warehouse_mode == "Zert":
                    self.inventur_data_zert = [
                        d for d in self.inventur_data_zert if d.get("charge") != charge]
                    self.not_found_data_zert = [
                        d for d in self.not_found_data_zert if d.get("charge") != charge]
                    self.save_zert_excel()
                elif self.warehouse_mode == "Granules":
                    key = _batch_key(charge)
                    self.inventur_data_granules = [
                        d for d in self.inventur_data_granules if _batch_key(d.get("charge")) != key]
                    self.not_found_data_granules = [
                        d for d in self.not_found_data_granules if _batch_key(d.get("charge")) != key]
                    self.save_granules_excel()
                else:
                    self.inventur_data = [
                        d for d in self.inventur_data if d.get("charge") != charge]
                    self.not_found_data = [
                        d for d in self.not_found_data if d.get("charge") != charge]
                    self.save_halb_excel()

                self.update_list()
                self.status_var.set(f"Entry deleted: {charge}")
                self.logger.info(f"Entry deleted: {charge}")

    # ------------------------------------------------------------------
    # Misc helpers
    # ------------------------------------------------------------------

    def toggle_fullscreen(self):
        current = self.root.attributes("-fullscreen")
        self.root.attributes("-fullscreen", not current)
        if not current:
            self.status_var.set("Fullscreen ON — press F11 to exit")
        else:
            self.status_var.set("Fullscreen OFF")

    def ensure_scan_focus(self, event=None):
        if hasattr(self, "scan_entry") and not self.current_scan:
            self.root.after(100, lambda: self.scan_entry.focus_set())

    def _manual_save(self):
        if self.warehouse_mode == "Zert":
            self.save_zert_excel()
        elif self.warehouse_mode == "KMAT":
            self.save_kmat_excel()
        elif self.warehouse_mode == "WIP":
            self.save_wip_excel()
        elif self.warehouse_mode == "Granules":
            self.save_granules_excel()
        else:
            self.save_halb_excel()
        self.status_var.set("Manually saved.")

    def quit_app(self):
        if messagebox.askyesno("Quit", "Are you sure you want to quit?"):
            self.logger.info("Application closed by user")
            self.root.quit()

    # ------------------------------------------------------------------
    # Run
    # ------------------------------------------------------------------

    def run(self):
        self.root.mainloop()


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def main():
    try:
        app = InventurApp()
        app.run()
    except Exception as e:
        messagebox.showerror("Critical Error", f"Unexpected error:\n{e}")
        logging.error(f"Critical error: {e}")
        raise


if __name__ == "__main__":
    main()
